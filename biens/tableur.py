"""Import et export du référentiel des biens au format Excel (.xlsx) ou CSV.

Une ligne = un lot. Le lot est identifié par sa référence ICS : un nouvel
import met à jour les lots existants et crée les autres. Une cellule vide ne
remplace jamais une valeur déjà saisie dans l'application.
"""

import csv
import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Bailleur, Bien


@dataclass(frozen=True)
class Colonne:
    entete: str
    champ: str
    type: str = "texte"  # texte, choix, booleen, decimal, entier, date
    modele: type = Bien
    obligatoire: bool = False
    largeur: int = 18


COLONNES = [
    Colonne("Réf. ICS lot", "ref_ics", obligatoire=True, largeur=14),
    Colonne("Réf. ICS bailleur", "ref_ics", modele=Bailleur, largeur=14),
    Colonne("Type de bailleur", "type", "choix", modele=Bailleur, largeur=20),
    Colonne("Bailleur (nom ou raison sociale)", "nom", modele=Bailleur, obligatoire=True, largeur=28),
    Colonne("Bailleur prénom", "prenom", modele=Bailleur),
    Colonne("Bailleur adresse", "adresse", modele=Bailleur, largeur=30),
    Colonne("Bailleur code postal", "code_postal", modele=Bailleur, largeur=12),
    Colonne("Bailleur ville", "ville", modele=Bailleur),
    Colonne("Bailleur e-mail", "email", modele=Bailleur, largeur=24),
    Colonne("Usage", "usage", "choix", largeur=20),
    Colonne("Meublé", "meuble", "booleen", largeur=10),
    Colonne("Adresse", "adresse", obligatoire=True, largeur=30),
    Colonne("Complément", "complement", largeur=24),
    Colonne("Code postal", "code_postal", obligatoire=True, largeur=12),
    Colonne("Ville", "ville", obligatoire=True),
    Colonne("Étage", "etage", largeur=10),
    Colonne("N° appartement", "numero_appartement", largeur=12),
    Colonne("N° parking", "numero_parking", largeur=12),
    Colonne("N° cellier", "numero_cellier", largeur=12),
    Colonne("Identifiant fiscal", "identifiant_fiscal", largeur=16),
    Colonne("Syndic", "syndic", largeur=20),
    Colonne("Type d'habitat", "type_habitat", "choix"),
    Colonne("Régime juridique", "regime_juridique", "choix"),
    Colonne("Période de construction", "periode_construction", "choix", largeur=22),
    Colonne("Surface (m²)", "surface", "decimal", largeur=12),
    Colonne("Pièces principales", "nb_pieces", "entier", largeur=12),
    Colonne("Équipements", "equipements", largeur=30),
    Colonne("Chauffage", "chauffage", "choix", largeur=12),
    Colonne("Eau chaude", "eau_chaude", "choix", largeur=12),
    Colonne("Accès TIC", "acces_tic", largeur=20),
    Colonne("Annexes privatives", "annexes", largeur=24),
    Colonne("Parties communes", "parties_communes", largeur=24),
    Colonne("Détecteur de fumée", "detecteur_fumee", "booleen", largeur=10),
    Colonne("Classe DPE", "classe_dpe", "choix", largeur=10),
    Colonne("Énergie min (€/an)", "depenses_energie_min", "entier", largeur=12),
    Colonne("Énergie max (€/an)", "depenses_energie_max", "entier", largeur=12),
    Colonne("Date DPE", "date_dpe", "date", largeur=12),
    Colonne("Zone tendue", "zone_tendue", "booleen", largeur=10),
    Colonne("Loyer de référence majoré (€/m²)", "loyer_reference_majore", "decimal", largeur=14),
    Colonne("Dernier loyer HC (€)", "dernier_loyer", "decimal", largeur=12),
    Colonne("Charges (€)", "charges", "decimal", largeur=12),
    Colonne("Dispositif fiscal", "dispositif_fiscal", largeur=20),
    Colonne("Observations", "observations", largeur=30),
]

VRAI = {"oui", "o", "x", "1", "vrai", "true", "yes"}
FAUX = {"non", "n", "0", "faux", "false", "no"}


def normaliser(texte):
    """« Réf. ICS lot » et « ref ics lot » désignent la même colonne."""
    texte = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", texte.lower()).strip()


def choix_de(colonne):
    return colonne.modele._meta.get_field(colonne.champ).choices


class ErreurCellule(Exception):
    pass


def convertir(colonne, brut):
    """Convertit une cellule ; renvoie None pour une cellule vide."""
    if brut is None or (isinstance(brut, str) and not brut.strip()):
        return None
    if colonne.type == "texte":
        if isinstance(brut, float) and brut.is_integer():
            brut = int(brut)  # référence ou code postal lu comme nombre
        return str(brut).strip()
    if colonne.type == "choix":
        cle = normaliser(brut)
        for valeur, libelle in choix_de(colonne):
            if cle in (normaliser(valeur), normaliser(libelle)):
                return valeur
        permis = ", ".join(libelle for _, libelle in choix_de(colonne))
        raise ErreurCellule(f"« {brut} » n'est pas une valeur permise ({permis})")
    if colonne.type == "booleen":
        if isinstance(brut, bool):
            return brut
        cle = normaliser(brut)
        if cle in VRAI:
            return True
        if cle in FAUX:
            return False
        raise ErreurCellule(f"« {brut} » : indiquer oui ou non")
    if colonne.type in ("decimal", "entier"):
        if isinstance(brut, (int, float, Decimal)) and not isinstance(brut, bool):
            nombre = Decimal(str(brut))
        else:
            texte = str(brut).replace(" ", "").replace(" ", "").replace("€", "")
            try:
                nombre = Decimal(texte.replace(",", "."))
            except InvalidOperation:
                raise ErreurCellule(f"« {brut} » n'est pas un nombre") from None
        if colonne.type == "entier":
            if nombre != nombre.to_integral_value():
                raise ErreurCellule(f"« {brut} » doit être un nombre entier")
            return int(nombre)
        return nombre
    if colonne.type == "date":
        if isinstance(brut, dt.datetime):
            return brut.date()
        if isinstance(brut, dt.date):
            return brut
        for format_ in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(str(brut).strip(), format_).date()
            except ValueError:
                pass
        raise ErreurCellule(f"« {brut} » n'est pas une date (JJ/MM/AAAA)")
    raise ValueError(colonne.type)


# --- Lecture -------------------------------------------------------------------


class PointVirgule(csv.excel):
    """Séparateur par défaut d'un CSV enregistré par Excel en français."""

    delimiter = ";"


def lire_lignes(nom_fichier, contenu):
    """Renvoie les lignes du fichier sous forme de listes de cellules."""
    if nom_fichier.lower().endswith((".xlsx", ".xlsm")):
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        feuille = classeur.worksheets[0]
        return [list(ligne) for ligne in feuille.iter_rows(values_only=True)]
    if nom_fichier.lower().endswith(".csv"):
        for encodage in ("utf-8-sig", "cp1252"):
            try:
                texte = contenu.decode(encodage)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t")
        except csv.Error:
            dialecte = PointVirgule
        return [ligne for ligne in csv.reader(io.StringIO(texte), dialecte)]
    raise ValueError("Format non pris en charge : fournir un fichier .xlsx ou .csv.")


@dataclass
class Rapport:
    crees: int = 0
    mis_a_jour: int = 0
    bailleurs_crees: int = 0
    erreurs: list = field(default_factory=list)  # (n° de ligne, message)
    colonnes_ignorees: list = field(default_factory=list)

    @property
    def total(self):
        return self.crees + self.mis_a_jour


class AnnulerSimulation(Exception):
    pass


def importer(nom_fichier, contenu, simulation=False):
    """Importe le fichier. En simulation, tout est vérifié puis annulé."""
    rapport = Rapport()
    lignes = lire_lignes(nom_fichier, contenu)
    if not lignes:
        rapport.erreurs.append((1, "Le fichier est vide."))
        return rapport

    par_entete = {normaliser(c.entete): c for c in COLONNES}
    index = {}
    for position, entete in enumerate(lignes[0]):
        if entete is None or not str(entete).strip():
            continue
        colonne = par_entete.get(normaliser(entete))
        if colonne:
            index[colonne] = position
        else:
            rapport.colonnes_ignorees.append(str(entete))
    manquantes = [c.entete for c in COLONNES if c.obligatoire and c not in index]
    if manquantes:
        rapport.erreurs.append((1, "Colonnes obligatoires absentes : " + ", ".join(manquantes)))
        return rapport

    try:
        with transaction.atomic():
            for numero, cellules in enumerate(lignes[1:], start=2):
                if all(c is None or str(c).strip() == "" for c in cellules):
                    continue
                try:
                    with transaction.atomic():
                        _importer_ligne(cellules, index, rapport)
                except (ErreurCellule, ValidationError) as erreur:
                    rapport.erreurs.append((numero, _message(erreur)))
            if simulation:
                raise AnnulerSimulation
    except AnnulerSimulation:
        pass
    return rapport


def _message(erreur):
    if isinstance(erreur, ValidationError) and hasattr(erreur, "message_dict"):
        return " ; ".join(
            f"{champ} : {' '.join(messages)}" for champ, messages in erreur.message_dict.items()
        )
    if isinstance(erreur, ValidationError):
        return " ".join(erreur.messages)
    return str(erreur)


def _importer_ligne(cellules, index, rapport):
    valeurs = {Bailleur: {}, Bien: {}}
    for colonne, position in index.items():
        brut = cellules[position] if position < len(cellules) else None
        try:
            valeur = convertir(colonne, brut)
        except ErreurCellule as erreur:
            raise ErreurCellule(f"{colonne.entete} : {erreur}") from None
        if valeur is not None:
            valeurs[colonne.modele][colonne.champ] = valeur
        elif colonne.obligatoire:
            raise ErreurCellule(f"{colonne.entete} : valeur obligatoire")

    bailleur = _trouver_bailleur(valeurs[Bailleur])
    if bailleur is None:
        bailleur = Bailleur()
        rapport.bailleurs_crees += 1
    for champ, valeur in valeurs[Bailleur].items():
        setattr(bailleur, champ, valeur)
    bailleur.full_clean()
    bailleur.save()

    ref = valeurs[Bien]["ref_ics"]
    bien = Bien.objects.filter(ref_ics=ref).first()
    nouveau = bien is None
    if nouveau:
        bien = Bien()
    bien.bailleur = bailleur
    for champ, valeur in valeurs[Bien].items():
        setattr(bien, champ, valeur)
    bien.full_clean()
    bien.save()
    if nouveau:
        rapport.crees += 1
    else:
        rapport.mis_a_jour += 1


def _trouver_bailleur(valeurs):
    if valeurs.get("ref_ics"):
        trouve = Bailleur.objects.filter(ref_ics=valeurs["ref_ics"]).first()
        if trouve:
            return trouve
    return Bailleur.objects.filter(
        nom__iexact=valeurs["nom"], prenom__iexact=valeurs.get("prenom", "")
    ).first()


# --- Écriture ------------------------------------------------------------------


def _valeur_export(colonne, objet):
    valeur = getattr(objet, colonne.champ)
    if colonne.type == "choix":
        return dict(choix_de(colonne)).get(valeur, "") if valeur else None
    if colonne.type == "booleen":
        return "oui" if valeur else "non"
    if isinstance(valeur, Decimal):
        return float(valeur)
    return valeur if valeur not in ("", None) else None


def classeur(biens=()):
    """Classeur au format d'import : modèle vide ou export des biens donnés."""
    wb = Workbook()
    feuille = wb.active
    feuille.title = "Biens"
    feuille.append([c.entete for c in COLONNES])
    for position, colonne in enumerate(COLONNES, start=1):
        cellule = feuille.cell(row=1, column=position)
        cellule.font = Font(bold=True)
        if colonne.obligatoire:
            cellule.fill = PatternFill("solid", fgColor="FDE9C9")
        feuille.column_dimensions[get_column_letter(position)].width = colonne.largeur
    feuille.freeze_panes = "B2"

    for bien in biens:
        feuille.append(
            [
                _valeur_export(c, bien.bailleur if c.modele is Bailleur else bien)
                for c in COLONNES
            ]
        )

    # Listes déroulantes pour les colonnes à valeurs fermées.
    aide = wb.create_sheet("Valeurs permises")
    aide.append(["Colonne", "Valeurs"])
    aide["A1"].font = aide["B1"].font = Font(bold=True)
    aide.column_dimensions["A"].width = 26
    aide.column_dimensions["B"].width = 80
    for position, colonne in enumerate(COLONNES, start=1):
        if colonne.type == "choix":
            libelles = [libelle for _, libelle in choix_de(colonne)]
        elif colonne.type == "booleen":
            libelles = ["oui", "non"]
        else:
            continue
        aide.append([colonne.entete, ", ".join(libelles)])
        liste = '"' + ",".join(libelles) + '"'
        if len(liste) <= 255:
            validation = DataValidation(type="list", formula1=liste, allow_blank=True)
            lettre = get_column_letter(position)
            validation.add(f"{lettre}2:{lettre}2000")
            feuille.add_data_validation(validation)
    aide.append([])
    aide.append(["Obligatoire", "Colonnes surlignées : " + ", ".join(
        c.entete for c in COLONNES if c.obligatoire
    )])
    aide.append(["Mise à jour", "Un lot déjà connu (même Réf. ICS lot) est mis à jour ; "
                 "une cellule vide ne remplace pas la valeur existante."])

    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
