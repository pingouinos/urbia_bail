import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from . import tableur
from .models import Bailleur, Bien

ENTETES = [
    "Réf. ICS lot", "Réf. ICS bailleur", "Type de bailleur", "Bailleur (nom ou raison sociale)",
    "Bailleur prénom", "Usage", "Adresse", "Complément", "Code postal", "Ville",
    "Surface (m²)", "Pièces principales", "Classe DPE", "Zone tendue", "Date DPE",
]


def ligne(ref="L001", **autres):
    valeurs = {
        "Réf. ICS lot": ref, "Réf. ICS bailleur": "B01", "Type de bailleur": "Personne physique",
        "Bailleur (nom ou raison sociale)": "Durand", "Bailleur prénom": "Anne",
        "Usage": "Habitation", "Adresse": "12 rue des Lois", "Complément": "Bât. A, 2e étage",
        "Code postal": "31000", "Ville": "Toulouse", "Surface (m²)": 45.5,
        "Pièces principales": 2, "Classe DPE": "D", "Zone tendue": "oui", "Date DPE": "15/03/2024",
    }
    valeurs.update(autres)
    return [valeurs.get(entete) for entete in ENTETES]


def xlsx(*lignes, entetes=ENTETES):
    wb = Workbook()
    wb.active.append(entetes)
    for valeurs in lignes:
        wb.active.append(valeurs)
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


class ImportTests(TestCase):
    def test_import_cree_bien_et_bailleur(self):
        rapport = tableur.importer("biens.xlsx", xlsx(ligne()))
        self.assertEqual((rapport.crees, rapport.bailleurs_crees, rapport.erreurs), (1, 1, []))
        bien = Bien.objects.get(ref_ics="L001")
        self.assertEqual(bien.surface, Decimal("45.50"))
        self.assertEqual(bien.classe_dpe, "D")
        self.assertTrue(bien.zone_tendue)
        self.assertEqual(str(bien.date_dpe), "2024-03-15")
        self.assertEqual(bien.bailleur.ref_ics, "B01")

    def test_reimport_met_a_jour_sans_effacer_les_cellules_vides(self):
        tableur.importer("biens.xlsx", xlsx(ligne()))
        rapport = tableur.importer("biens.xlsx", xlsx(ligne(**{"Surface (m²)": None, "Classe DPE": "C"})))
        self.assertEqual((rapport.crees, rapport.mis_a_jour, rapport.bailleurs_crees), (0, 1, 0))
        bien = Bien.objects.get(ref_ics="L001")
        self.assertEqual(bien.surface, Decimal("45.50"))
        self.assertEqual(bien.classe_dpe, "C")
        self.assertEqual(Bailleur.objects.count(), 1)

    def test_erreurs_signalees_par_ligne_sans_bloquer_les_autres(self):
        rapport = tableur.importer("biens.xlsx", xlsx(
            ligne("L001"),
            ligne("L002", **{"Classe DPE": "Z"}),
            ligne("L003", **{"Surface (m²)": "quarante"}),
            ligne("L004", **{"Ville": None}),
        ))
        self.assertEqual(rapport.crees, 1)
        self.assertEqual([numero for numero, _ in rapport.erreurs], [3, 4, 5])
        self.assertIn("Classe DPE", rapport.erreurs[0][1])
        self.assertIn("nombre", rapport.erreurs[1][1])
        self.assertIn("obligatoire", rapport.erreurs[2][1])
        self.assertEqual(list(Bien.objects.values_list("ref_ics", flat=True)), ["L001"])

    def test_simulation_n_enregistre_rien(self):
        rapport = tableur.importer("biens.xlsx", xlsx(ligne()), simulation=True)
        self.assertEqual(rapport.crees, 1)
        self.assertFalse(Bien.objects.exists())
        self.assertFalse(Bailleur.objects.exists())

    def test_colonne_obligatoire_absente(self):
        rapport = tableur.importer("biens.xlsx", xlsx(["X"], entetes=["Adresse"]))
        self.assertIn("Réf. ICS lot", rapport.erreurs[0][1])

    def test_csv_excel_francais(self):
        contenu = (
            "Ref ICS lot;Bailleur (nom ou raison sociale);Adresse;Code postal;Ville;Surface (m²);Usage;Meublé\n"
            "L010;Société Générale Immo;5 allée Jean Jaurès;31000;Toulouse;32,5;Local professionnel;non\n"
        ).encode("cp1252")
        rapport = tableur.importer("export.csv", contenu)
        self.assertEqual(rapport.erreurs, [])
        bien = Bien.objects.get(ref_ics="L010")
        self.assertEqual(bien.surface, Decimal("32.50"))
        self.assertEqual(bien.usage, Bien.Usage.PROFESSIONNEL)
        self.assertEqual(bien.bailleur.nom, "Société Générale Immo")

    def test_bailleur_retrouve_par_nom_sans_reference(self):
        Bailleur.objects.create(nom="Martin", prenom="Paul")
        tableur.importer("biens.xlsx", xlsx(ligne(**{
            "Réf. ICS bailleur": None, "Bailleur (nom ou raison sociale)": "MARTIN", "Bailleur prénom": "paul",
        })))
        self.assertEqual(Bailleur.objects.count(), 1)

    def test_export_reimportable(self):
        tableur.importer("biens.xlsx", xlsx(ligne("L001"), ligne("L002")))
        contenu = tableur.classeur(Bien.objects.select_related("bailleur"))
        feuille = load_workbook(io.BytesIO(contenu)).active
        self.assertEqual(feuille.max_row, 3)
        Bien.objects.update(surface=None)
        rapport = tableur.importer("export.xlsx", contenu)
        self.assertEqual((rapport.mis_a_jour, rapport.erreurs), (2, []))
        self.assertEqual(Bien.objects.get(ref_ics="L001").surface, Decimal("45.50"))

    def test_modele_vide_a_listes_deroulantes(self):
        feuille = load_workbook(io.BytesIO(tableur.classeur())).active
        self.assertEqual(feuille.max_row, 1)
        self.assertTrue(feuille.data_validations.dataValidation)


class BienTests(TestCase):
    def test_logement_g_signale(self):
        bailleur = Bailleur.objects.create(nom="Durand")
        bien = Bien(bailleur=bailleur, adresse="1 rue", code_postal="31000", ville="Toulouse", classe_dpe="G")
        self.assertTrue(bien.location_interdite)
        bien.usage = Bien.Usage.COMMERCIAL
        self.assertFalse(bien.location_interdite)

    def test_duree_bail_selon_bailleur(self):
        self.assertEqual(Bailleur(type=Bailleur.Type.SCI_FAMILIALE).duree_bail_nu_ans, 3)
        self.assertEqual(Bailleur(type=Bailleur.Type.PERSONNE_MORALE).duree_bail_nu_ans, 6)


@override_settings(MFA_OBLIGATOIRE=False)
class VuesTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_user("gestion"))

    def test_parcours_import_verification_puis_confirmation(self):
        fichier = SimpleUploadedFile("biens.xlsx", xlsx(ligne("L001"), ligne("L002", **{"Classe DPE": "Z"})))
        reponse = self.client.post(reverse("biens:importer"), {"fichier": fichier})
        self.assertContains(reponse, "1 bien(s) à créer")
        self.assertContains(reponse, "Classe DPE")
        self.assertFalse(Bien.objects.exists())

        reponse = self.client.post(reverse("biens:importer_confirmer"), follow=True)
        self.assertContains(reponse, "1 bien(s) créé(s)")
        self.assertTrue(Bien.objects.filter(ref_ics="L001").exists())

        # Le fichier n'est importé qu'une fois.
        reponse = self.client.post(reverse("biens:importer_confirmer"), follow=True)
        self.assertContains(reponse, "Aucun fichier en attente")

    def test_fichier_illisible(self):
        fichier = SimpleUploadedFile("biens.xlsx", b"pas un classeur")
        reponse = self.client.post(reverse("biens:importer"), {"fichier": fichier})
        self.assertContains(reponse, "Impossible de lire ce fichier")

    def test_liste_recherche_et_export(self):
        tableur.importer("biens.xlsx", xlsx(ligne("L001"), ligne("L002", Ville="Blagnac", **{"Code postal": "31700"})))
        reponse = self.client.get(reverse("biens:liste"), {"q": "blagnac"})
        self.assertContains(reponse, "1 bien(s)")
        reponse = self.client.get(reverse("biens:exporter"), {"q": "blagnac"})
        self.assertEqual(load_workbook(io.BytesIO(reponse.content)).active.max_row, 2)

    def test_creation_bien_et_bailleur_sans_reference_ics(self):
        for nom in ("Durand", "Martin"):
            reponse = self.client.post(reverse("biens:bailleur_creer"), {"type": "physique", "nom": nom})
            self.assertEqual(reponse.status_code, 302)
        self.assertEqual(Bailleur.objects.filter(ref_ics__isnull=True).count(), 2)

        bailleur = Bailleur.objects.first()
        reponse = self.client.post(reverse("biens:bien_creer"), {
            "bailleur": bailleur.pk, "usage": "habitation", "actif": "on",
            "adresse": "3 place du Capitole", "code_postal": "31000", "ville": "Toulouse",
        })
        bien = Bien.objects.get()
        self.assertRedirects(reponse, bien.get_absolute_url())
        reponse = self.client.get(bien.get_absolute_url())
        self.assertContains(reponse, "À compléter avant de rédiger un bail")
        self.assertContains(reponse, "gestion")  # auteur dans l'historique

    def test_depenses_energie_incoherentes(self):
        bailleur = Bailleur.objects.create(nom="Durand")
        reponse = self.client.post(reverse("biens:bien_creer"), {
            "bailleur": bailleur.pk, "usage": "habitation", "adresse": "1 rue",
            "code_postal": "31000", "ville": "Toulouse",
            "depenses_energie_min": 900, "depenses_energie_max": 500,
        })
        self.assertContains(reponse, "Le maximum doit être supérieur au minimum")

    def test_toutes_les_pages_s_affichent(self):
        tableur.importer("biens.xlsx", xlsx(ligne("L001")))
        bien, bailleur = Bien.objects.get(), Bailleur.objects.get()
        for url in [
            reverse("biens:liste"), reverse("biens:bien", args=[bien.pk]),
            reverse("biens:bien_creer"), reverse("biens:bien_modifier", args=[bien.pk]),
            reverse("biens:bailleurs"), reverse("biens:bailleur", args=[bailleur.pk]),
            reverse("biens:bailleur_creer"), reverse("biens:bailleur_modifier", args=[bailleur.pk]),
            reverse("biens:importer"), reverse("biens:modele"), reverse("accueil"),
        ]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
