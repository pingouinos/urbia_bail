import io
import shutil
import tempfile
from decimal import Decimal
from unittest import mock

from core.testing import creer_gestionnaire, remplir_agence

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from openpyxl import load_workbook

from proprietes.models import Bailleur, Propriete
from comptes.roles import ADMINISTRATEURS, attribuer_role
from documents import generation
from documents.generation import ConversionImpossible

from .models import Mandat


def texte_docx(contenu):
    return "\n".join(p.text for p in Document(io.BytesIO(contenu)).paragraphs)


class Donnees:
    def creer_propriete(self, **autres):
        bailleur = Bailleur.objects.create(
            civilite="M.", nom="Durand", prenom="Paul", conjoint_civilite="Mme",
            conjoint_nom="Martin", conjoint_prenom="Anne",
            adresse="9 rue des Lilas", code_postal="31270", ville="Frouzins",
        )
        valeurs = {
            "bailleur": bailleur, "adresse": "12 avenue de Muret", "code_postal": "31300",
            "ville": "Toulouse", "surface": Decimal("39"), "nb_pieces": 2, "meuble": True,
            "zone_tendue": True, "dernier_loyer": Decimal("854"), "charges": Decimal("45"),
        }
        valeurs.update(autres)
        return Propriete.objects.create(**valeurs)

    def creer_mandat(self, propriete=None, **autres):
        remplir_agence()
        autres.setdefault("conclusion", Mandat.Conclusion.AGENCE)
        mandat = Mandat.depuis_propriete(propriete or self.creer_propriete())
        for champ, valeur in autres.items():
            setattr(mandat, champ, valeur)
        mandat.save()
        return mandat


class MandatTests(Donnees, TestCase):
    def test_numeros_du_registre_se_suivent(self):
        propriete = self.creer_propriete()
        numeros = [self.creer_mandat(propriete).numero for _ in range(3)]
        self.assertEqual(numeros, [1, 2, 3])

    def test_prerempli_depuis_les_fiches(self):
        mandat = Mandat.depuis_propriete(self.creer_propriete())
        self.assertEqual(mandat.mandant, "Monsieur DURAND Paul et Madame MARTIN Anne")
        self.assertEqual(mandat.mandant_code_postal_ville, "31270 Frouzins")
        self.assertEqual(mandat.designation, "Appartement T2 meublé de 39 m²")
        self.assertEqual(mandat.adresse_bien, "12 avenue de Muret 31300 TOULOUSE")
        self.assertEqual((mandat.loyer_cc, mandat.charges), (Decimal("899"), Decimal("45")))

    def test_designation_personne_morale(self):
        self.assertEqual(Bailleur(type=Bailleur.Type.SCI_FAMILIALE, nom="SCI Les Pins").designation, "SCI Les Pins")

    def test_duree(self):
        self.assertEqual(Mandat(duree_ans=3).duree_texte, "une durée de trois ans")
        self.assertEqual(
            Mandat(duree_ans=1, reconductions=3).duree_texte,
            "une durée d’un an, le contrat se renouvellera par tacite reconduction à trois reprises",
        )

    def test_honoraires_conformes(self):
        mandat = self.creer_mandat(
            honoraires_bail_bailleur=390, honoraires_bail_locataire=390,
            honoraires_edl_bailleur=117, honoraires_edl_locataire=117,
        )
        self.assertEqual(mandat.alertes, [])

    def test_part_locataire_superieure_a_celle_du_bailleur(self):
        mandat = self.creer_mandat(
            honoraires_bail_bailleur=0, honoraires_bail_locataire=390, honoraires_edl_bailleur=0, honoraires_edl_locataire=117,
        )
        self.assertEqual(len(mandat.alertes), 2)
        self.assertIn("dépasse celle du bailleur", mandat.alertes[0])

    def test_plafond_par_m2(self):
        mandat = self.creer_mandat(honoraires_bail_bailleur=400, honoraires_bail_locataire=400)
        self.assertEqual(len(mandat.alertes), 1)
        self.assertIn("plafond légal de 390 € pour 39 m²", mandat.alertes[0])
        mandat.propriete.zone_tendue = False  # 8 €/m² hors zone tendue
        mandat.honoraires_bail_bailleur = mandat.honoraires_bail_locataire = Decimal("312")
        self.assertEqual(mandat.alertes, [])

    def test_stationnement_non_soumis(self):
        propriete = self.creer_propriete(usage=Propriete.Usage.STATIONNEMENT, surface=None)
        self.assertEqual(self.creer_mandat(propriete, honoraires_bail_locataire=50).alertes, [])

    def test_document_word(self):
        mandat = self.creer_mandat(
            date_signature="2026-07-20", assurance_loyers_impayes=True,
            base_gestion=Mandat.BaseGestion.HORS_CHARGES,
            honoraires_bail_bailleur=390, honoraires_bail_locataire=390,
            honoraires_edl_bailleur=Decimal("58.50"), honoraires_edl_locataire=Decimal("58.50"),
        )
        mandat.refresh_from_db()
        texte = texte_docx(generation.generer_docx("mandat_gestion.docx", mandat.contexte_document()))
        self.assertNotIn("{{", texte)
        self.assertIn("MANDAT DE GESTION N° 1", texte)
        self.assertIn("Monsieur DURAND Paul et Madame MARTIN Anne", texte)
        self.assertIn("Adresse du bien\xa0: 12 avenue de Muret 31300 TOULOUSE", texte)
        self.assertIn("Montant\xa0: 899€ Charges comprises dont 45€ de charges", texte)
        self.assertIn("pour une durée de trois ans.", texte)
        self.assertIn("☒Le mandant souhaite adhérer", texte)
        self.assertIn("à 7% TTC", texte)
        self.assertIn("encaissements mensuels (hors charges)", texte)
        self.assertIn("la somme de 780 euros TTC, soit 390 euros TTC à la charge du bailleur et 390 euros", texte)
        self.assertIn("la somme de 117 euros TTC, soit 58,50 euros TTC", texte)
        self.assertIn("Fait à Toulouse, en deux exemplaires originaux dont un remis au mandant, le : 20/07/2026", texte)

    def test_montant(self):
        self.assertEqual(generation.montant(Decimal("1234.5")), "1 234,50")
        self.assertEqual(generation.montant(Decimal("390.00")), "390")


@override_settings(MFA_OBLIGATOIRE=False)
class VuesTests(Donnees, TestCase):
    def setUp(self):
        self.dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.dossier)
        reglage = override_settings(MEDIA_ROOT=self.dossier)
        reglage.enable()
        self.addCleanup(reglage.disable)
        self.client.force_login(creer_gestionnaire("gestion"))

    def test_creation_depuis_le_propriete(self):
        propriete = self.creer_propriete()
        reponse = self.client.get(reverse("mandats:creer"))
        self.assertContains(reponse, str(propriete))
        reponse = self.client.get(reverse("mandats:creer"), {"propriete": propriete.pk})
        self.assertContains(reponse, "Monsieur DURAND Paul et Madame MARTIN Anne")
        donnees = {
            "numero": "412", "propriete": propriete.pk, "mandant": "Monsieur DURAND Paul", "mandant_adresse": "9 rue des Lilas",
            "mandant_code_postal_ville": "31270 Frouzins", "designation": "Appartement T2",
            "adresse_bien": "12 avenue de Muret", "loyer_cc": "899", "charges": "45",
            "duree_ans": 3, "reconductions": 0, "taux_gestion": "7", "base_gestion": "charges_comprises",
            "taux_assurance": "2.90", "honoraires_bail_bailleur": "0", "honoraires_bail_locataire": "390",
            "honoraires_edl_bailleur": "0", "honoraires_edl_locataire": "0",
            "conclusion": "agence", "lieu_signature": "Toulouse",
        }
        reponse = self.client.post(reverse("mandats:creer"), donnees, follow=True)
        mandat = Mandat.objects.get()
        # Le mandat rédigé s'ouvre sur sa page, à l'étape de la signature.
        self.assertRedirects(reponse, mandat.get_absolute_url() + "#etape-2")
        self.assertContains(reponse, "dépasse celle du bailleur")
        self.assertEqual(propriete.mandat_en_cours, mandat)
        # Le numéro est celui du registre papier, reporté à la main.
        self.assertEqual(mandat.numero, 412)
        autre = self.creer_propriete(adresse="5 rue du Taur")
        page = self.client.get(reverse("mandats:creer"), {"propriete": autre.pk})
        self.assertContains(page, "Le plus grand numéro déjà saisi ici est le 412.")
        reponse = self.client.post(reverse("mandats:creer"), {**donnees, "propriete": autre.pk})
        self.assertContains(reponse, "Ce numéro est déjà porté par un autre mandat.")
        reponse = self.client.post(reverse("mandats:creer"), {**donnees, "propriete": autre.pk, "numero": ""})
        self.assertIn("numero", reponse.context["form"].errors)
        self.assertEqual(Mandat.objects.count(), 1)
        # Une erreur de recopie se corrige, et la correction reste dans l'historique.
        self.client.post(mandat.get_absolute_url(), {**donnees, "numero": "413"})
        mandat.refresh_from_db()
        self.assertEqual((mandat.numero, mandat.history.count()), (413, 2))
        # Un bien sous mandat n'est plus proposé pour un nouveau mandat.
        self.assertNotContains(self.client.get(reverse("mandats:creer")), str(propriete))

    def test_telechargement_word_et_pdf(self):
        mandat = self.creer_mandat()
        reponse = self.client.get(reverse("mandats:word", args=[mandat.pk]))
        self.assertIn("Monsieur DURAND Paul", texte_docx(reponse.content))
        with mock.patch("documents.views.convertir_pdf", return_value=b"%PDF-1.7") as convertir:
            reponse = self.client.get(reverse("mandats:pdf", args=[mandat.pk]))
        self.assertEqual(reponse["Content-Type"], "application/pdf")
        self.assertEqual(reponse.content, b"%PDF-1.7")
        self.assertTrue(convertir.call_args.args[0].startswith(b"PK"))

    def test_pdf_indisponible(self):
        mandat = self.creer_mandat()
        with mock.patch("documents.views.convertir_pdf", side_effect=ConversionImpossible("hors service")):
            reponse = self.client.get(reverse("mandats:pdf", args=[mandat.pk]), follow=True)
        self.assertContains(reponse, "téléchargez la version Word")

    def test_conversion_gotenberg(self):
        reponse = mock.MagicMock()
        reponse.__enter__.return_value.read.return_value = b"%PDF"
        with mock.patch("urllib.request.urlopen", return_value=reponse) as urlopen:
            self.assertEqual(generation.convertir_pdf(b"docx", "m.docx"), b"%PDF")
        requete = urlopen.call_args.args[0]
        self.assertTrue(requete.full_url.endswith("/forms/libreoffice/convert"))
        self.assertIn(b'filename="m.docx"', requete.data)

    def test_conversion_impossible_dit_pourquoi(self):
        import urllib.error

        refus = urllib.error.URLError(ConnectionRefusedError(111, "Connection refused"))
        with mock.patch("urllib.request.urlopen", side_effect=refus), self.assertLogs("documents", "WARNING"):
            with self.assertRaisesMessage(ConversionImpossible, "arrêté ou en cours de redémarrage"):
                generation.convertir_pdf(b"docx")
        # LibreOffice pas encore prêt : une seconde tentative suffit.
        pas_pret = urllib.error.HTTPError("http://gotenberg", 503, "Service Unavailable", {}, io.BytesIO(b""))
        pret = mock.MagicMock()
        pret.__enter__.return_value.read.return_value = b"%PDF"
        with mock.patch("urllib.request.urlopen", side_effect=[pas_pret, pret]), mock.patch("time.sleep"):
            self.assertEqual(generation.convertir_pdf(b"docx"), b"%PDF")
        erreur = urllib.error.HTTPError("http://gotenberg", 500, "Internal", {}, io.BytesIO(b"libreoffice a plante"))
        with mock.patch("urllib.request.urlopen", side_effect=erreur), self.assertLogs("documents", "WARNING") as journal:
            with self.assertRaisesMessage(ConversionImpossible, "erreur 500"):
                generation.convertir_pdf(b"docx")
        self.assertIn("libreoffice a plante", journal.output[0])
        mandat = self.creer_mandat()
        with mock.patch("documents.views.convertir_pdf", side_effect=ConversionImpossible("le service ne répond pas")):
            reponse = self.client.get(reverse("mandats:pdf", args=[mandat.pk]), follow=True)
        self.assertContains(reponse, "indisponible (le service ne répond pas)")

    def test_signature_du_mandat(self):
        """La date de signature ne se saisit pas à la rédaction : elle
        s'enregistre une fois le mandat signé, avec ou sans le scan."""
        propriete = self.creer_propriete()
        page = self.client.get(reverse("mandats:creer"), {"propriete": propriete.pk})
        self.assertNotIn("date_signature", page.context["form"].fields)
        mandat = self.creer_mandat(propriete)
        self.assertContains(self.client.get(mandat.get_absolute_url()), "À signer")
        deposer = reverse("mandats:deposer_signe", args=[mandat.pk])
        demain = timezone.localdate() + timezone.timedelta(days=1)
        reponse = self.client.post(deposer, {"date_signature": demain.isoformat()}, follow=True)
        self.assertContains(reponse, "ne peut pas être dans le futur")
        # La date seule suffit ; le scan suit quand il est disponible.
        self.client.post(reverse("documents:preparer", args=["mandat", mandat.pk]))
        reponse = self.client.post(deposer, {"date_signature": "2026-09-15"}, follow=True)
        self.assertContains(reponse, "Signé le 15/09/2026")
        mandat.refresh_from_db()
        self.assertEqual((str(mandat.date_signature), bool(mandat.fichier_signe)), ("2026-09-15", False))
        fichier = SimpleUploadedFile("scan.pdf", b"%PDF-1.4 signe", content_type="application/pdf")
        self.client.post(deposer, {"date_signature": "2026-09-15", "fichier_signe": fichier})
        reponse = self.client.get(reverse("mandats:fichier_signe", args=[mandat.pk]))
        self.assertEqual(b"".join(reponse.streaming_content), b"%PDF-1.4 signe")
        reponse = self.client.post(
            deposer, {"date_signature": "2026-09-15", "fichier_signe": SimpleUploadedFile("virus.exe", b"MZ")}, follow=True,
        )
        self.assertContains(reponse, "ni un PDF ni une image")
        # L'extension ne suffit pas : un exécutable renommé en .pdf est refusé, sur le mandat comme sur le bail.
        faux = SimpleUploadedFile("evil.pdf", b"MZ\x90\x00", content_type="application/pdf")
        reponse = self.client.post(deposer, {"date_signature": "2026-09-15", "fichier_signe": faux}, follow=True)
        self.assertContains(reponse, "ni un PDF ni une image")
        mandat.refresh_from_db()
        # Le fichier est rangé sous un nom tiré au hasard et servi sans laisser le navigateur deviner son type.
        self.assertRegex(mandat.fichier_signe.name, r"^mandats/[0-9a-f]{32}\.pdf$")
        reponse = self.client.get(reverse("mandats:fichier_signe", args=[mandat.pk]))
        self.assertEqual((reponse["Content-Type"], reponse["X-Content-Type-Options"]), ("application/pdf", "nosniff"))

    def test_honoraires_proposes_au_metre_carre(self):
        """Visite, dossier et bail : plafond par m² selon la zone, reporté sur
        la part du bailleur et celle du locataire."""
        tendue = Mandat.depuis_propriete(self.creer_propriete())
        self.assertEqual((tendue.honoraires_bail_bailleur, tendue.honoraires_bail_locataire), (Decimal("390.00"), Decimal("390.00")))
        ailleurs = Mandat.depuis_propriete(self.creer_propriete(zone_tendue=False, surface=Decimal("52.5")))
        self.assertEqual((ailleurs.honoraires_bail_bailleur, ailleurs.honoraires_bail_locataire), (Decimal("420.00"), Decimal("420.00")))
        # L'état des lieux : 3 €/m² quelle que soit la zone.
        self.assertEqual((tendue.honoraires_edl_bailleur, tendue.honoraires_edl_locataire), (Decimal("117.00"), Decimal("117.00")))
        self.assertEqual((ailleurs.honoraires_edl_bailleur, ailleurs.honoraires_edl_locataire), (Decimal("157.50"), Decimal("157.50")))
        self.assertEqual(ailleurs.alertes, ailleurs.interdictions)
        # Sans surface, ou pour un parking, rien n'est proposé.
        self.assertEqual(Mandat.depuis_propriete(self.creer_propriete(surface=None)).honoraires_bail_locataire, 0)
        page = self.client.get(reverse("mandats:creer"), {"propriete": tendue.propriete.pk})
        self.assertContains(page, "39 m² × 10 €/m² (zone tendue) = 390 €")
        self.assertContains(page, "39 m² × 3 €/m² = 117 €")
        self.assertContains(page, 'name="honoraires_edl_bailleur" value="117.00"')
        self.assertContains(page, 'name="honoraires_bail_locataire" value="390.00"')

    def test_pieces_du_dossier(self):
        """Les pièces remises par le mandant se déposent depuis l'étape
        « Dossier » de la page du mandat, à tout moment : identité et RIB sur
        le bailleur, le reste sur la propriété."""
        from proprietes.models import Piece

        propriete = self.creer_propriete(regime_juridique="copropriete")
        mandat = self.creer_mandat(propriete)
        page = self.client.get(mandat.get_absolute_url())
        self.assertContains(page, "9 pièces à fournir")
        self.assertEqual([ligne["libelle"] for ligne in mandat.dossier if not ligne["obligatoire"]], ["Bail en cours", "État des lieux"])
        retour = mandat.get_absolute_url() + "?etape=3#etape-3"

        def deposer(type_, cible, **autres):
            nom = "proprietes:bailleur_piece_deposer" if cible is propriete.bailleur else "proprietes:piece_deposer"
            fichier = SimpleUploadedFile(f"{type_}.pdf", b"%PDF-1.7", content_type="application/pdf")
            return self.client.post(reverse(nom, args=[cible.pk]), {"type": type_, "fichier": fichier, "retour": retour, **autres})

        self.assertRedirects(deposer("identite", propriete.bailleur), retour, fetch_redirect_response=False)
        deposer("rib", propriete.bailleur)
        deposer("dpe", propriete, date_document="2020-01-10")
        identite = Piece.objects.get(type="identite")
        self.assertEqual((identite.bailleur, identite.propriete), (propriete.bailleur, None))
        self.assertTrue(identite.fichier.name.startswith("bailleurs/"))
        page = self.client.get(mandat.get_absolute_url(), {"etape": 3})
        self.assertEqual(page.context["ouverte"], 3)
        # Le DPE déposé est périmé : il reste à fournir.
        self.assertContains(page, "7 pièces à fournir")
        self.assertContains(page, "depuis le 31/12/2024")
        # Les pièces du bailleur valent pour ses autres mandats.
        autre = self.creer_mandat(Propriete.objects.create(
            bailleur=propriete.bailleur, adresse="5 rue du Taur", code_postal="31000", ville="Toulouse",
            surface=Decimal("30"), dernier_loyer=Decimal("500"), regime_juridique="monopropriete",
        ))
        manquantes = autre.pieces_manquantes
        self.assertNotIn("Pièce d'identité", manquantes)
        self.assertNotIn("Relevé annuel des charges de copropriété", manquantes)
        self.assertIn("Dernière taxe foncière", manquantes)
        self.assertContains(self.client.get(propriete.bailleur.get_absolute_url()), "Relevé d&#x27;identité bancaire")
        # Une société fournit en plus son justificatif d'immatriculation (extrait RNE ou Kbis) ; ses statuts sont facultatifs.
        sci = Bailleur.objects.create(type="sci_familiale", nom="SCI ESGT", representant_nom="Garcia", representant_prenom="Elodie")
        societe = self.creer_mandat(Propriete.objects.create(
            bailleur=sci, adresse="9 rue Riquet", code_postal="31000", ville="Toulouse", dernier_loyer=Decimal("500"),
        ))
        lignes = {ligne["libelle"]: ligne for ligne in societe.dossier}
        self.assertTrue(lignes["Justificatif d'immatriculation de la société"]["obligatoire"])
        self.assertFalse(lignes["Statuts de la société"]["obligatoire"])
        self.assertIn("extrait RNE", lignes["Justificatif d'immatriculation de la société"]["note"])
        self.assertEqual(lignes["Pièce d'identité"]["note"], "du représentant, pour une société")
        self.assertNotIn("Justificatif d'immatriculation de la société", [ligne["libelle"] for ligne in mandat.dossier])
        # Une pièce d'identité ne se dépose pas sur une propriété, ni un retour vers un autre site.
        reponse = self.client.post(reverse("proprietes:piece_deposer", args=[propriete.pk]), {
            "type": "identite", "fichier": SimpleUploadedFile("id.pdf", b"%PDF-1.7"), "retour": "https://exemple.org/",
        })
        self.assertRedirects(reponse, propriete.get_absolute_url() + "?onglet=documents", fetch_redirect_response=False)
        self.assertEqual(Piece.objects.filter(type="identite").count(), 1)

    def test_logo_de_l_agence_sur_le_document(self):
        """Le logo déposé sur la fiche de l'agence prend la place de celui de
        l'en-tête du modèle, sans être déformé."""
        import re
        import struct
        import zipfile

        from core.models import Organisation

        mandat = self.creer_mandat()
        sans_logo = generation.generer_docx("mandat_gestion.docx", mandat.contexte_document())
        self.assertNotIn("word/media/logo-agence.png", zipfile.ZipFile(io.BytesIO(sans_logo)).namelist())
        # Le modèle livré porte une image d'emplacement marquée LOGO ; sans logo d'agence elle ne s'imprime pas.
        modele = zipfile.ZipFile(generation.MODELES_PAR_DEFAUT / "mandat_gestion.docx").read("word/header1.xml").decode()
        self.assertIn('descr="logo"', modele)
        self.assertNotIn('descr="logo"', zipfile.ZipFile(io.BytesIO(sans_logo)).read("word/header1.xml").decode())
        # Un PNG carré de 400 × 400 points, dans un cadre prévu pour un logo large.
        png = b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + struct.pack(">II", 400, 400) + b"\x08\x06\x00\x00\x00"
        self.assertEqual(generation.dimensions_image(png), ("png", 400, 400))
        organisation = Organisation.courante()
        organisation.logo.save("logo.png", SimpleUploadedFile("logo.png", png), save=True)
        archive = zipfile.ZipFile(io.BytesIO(generation.generer_docx("mandat_gestion.docx", mandat.contexte_document())))
        self.assertEqual(archive.read("word/media/logo-agence.png"), png)
        self.assertIn('Target="media/logo-agence.png"', archive.read("word/_rels/header1.xml.rels").decode())
        entete = archive.read("word/header1.xml").decode()
        cx, cy = map(int, re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"', entete).groups())
        self.assertEqual((cx, cy), (883920, 883920))
        # Les autres images du document ne sont pas touchées.
        self.assertIn("word/media/image8.png", archive.namelist())
        # Un format que Word n'affiche pas laisse le logo du modèle.
        organisation.logo.save("logo.webp", SimpleUploadedFile("logo.webp", b"RIFF0000WEBPVP8 "), save=True)
        archive = zipfile.ZipFile(io.BytesIO(generation.generer_docx("mandat_gestion.docx", mandat.contexte_document())))
        self.assertNotIn("word/media/logo-agence.png", archive.namelist())

    def test_registre_et_export(self):
        self.creer_mandat()
        self.creer_mandat(date_fin="2026-01-01")
        self.assertContains(self.client.get(reverse("mandats:registre"), {"etat": "en_cours"}), "1 mandat(s)")
        reponse = self.client.get(reverse("mandats:registre_export"))
        self.assertEqual(load_workbook(io.BytesIO(reponse.content)).active.max_row, 3)

    def test_pages(self):
        mandat = self.creer_mandat()
        for url in [
            reverse("mandats:registre"), reverse("mandats:mandat", args=[mandat.pk]),
            reverse("proprietes:propriete", args=[mandat.propriete.pk]),
            reverse("accueil"),
        ]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_modeles_reserves_aux_administrateurs(self):
        self.assertEqual(self.client.get(reverse("documents:modeles")).status_code, 403)

    def test_remplacement_du_modele(self):
        admin = User.objects.create_user("admin")
        attribuer_role(admin, ADMINISTRATEURS)
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("documents:modeles")), "mandant_adresse")

        modele = Document(generation.MODELES_PAR_DEFAUT / "mandat_gestion.docx")
        modele.add_paragraph("Ligne ajoutée par l'agence pour {{ mandant }}.")
        flux = io.BytesIO()
        modele.save(flux)
        url = reverse("documents:deposer", args=["mandat_gestion.docx"])
        self.client.post(url, {"fichier": SimpleUploadedFile("m.docx", flux.getvalue())})
        mandat = self.creer_mandat()
        texte = texte_docx(self.client.get(reverse("mandats:word", args=[mandat.pk])).content)
        self.assertIn("Ligne ajoutée par l'agence pour Monsieur DURAND", texte)

        modele.add_paragraph("{{ balise_inconnue }}")
        flux = io.BytesIO()
        modele.save(flux)
        reponse = self.client.post(url, {"fichier": SimpleUploadedFile("m.docx", flux.getvalue())}, follow=True)
        self.assertContains(reponse, "Balises inconnues : balise_inconnue")

        self.client.post(reverse("documents:retablir", args=["mandat_gestion.docx"]))
        texte = texte_docx(self.client.get(reverse("mandats:word", args=[mandat.pk])).content)
        self.assertNotIn("Ligne ajoutée", texte)


@override_settings(MFA_OBLIGATOIRE=False)
class EtapesTests(Donnees, TestCase):
    """La page d'un mandat, comme celle du bail : une section par étape, un
    crayon tant qu'une étape se modifie, un cadenas une fois le mandat signé,
    et les flèches ‹ › de la frise."""

    def setUp(self):
        self.client.force_login(creer_gestionnaire("gestion"))
        self.signe_le = timezone.localdate() - timezone.timedelta(days=30)

    def etapes(self, mandat, **parametres):
        page = self.client.get(mandat.get_absolute_url(), parametres)
        return page, {section["numero"]: section["mode"] for section in page.context["sections"]}

    def crayons(self, page):
        return [section["numero"] for section in page.context["sections"] if section["crayon"]]

    def fleches(self, page):
        return tuple(fleche and fleche["numero"] for fleche in (page.context["precedente"], page.context["suivante"]))

    def test_ancienne_adresse_modifier(self):
        mandat = self.creer_mandat()
        self.assertRedirects(self.client.get(reverse("mandats:modifier", args=[mandat.pk])), mandat.get_absolute_url())

    def test_mandat_a_signer(self):
        mandat = self.creer_mandat()
        page, modes = self.etapes(mandat)
        self.assertEqual(modes, {1: "grisee", 2: "ouverte", 3: "a-venir", 4: "a-venir", 5: "a-venir"})
        self.assertEqual((self.crayons(page), self.fleches(page)), ([1], (1, None)))
        self.assertContains(page, "Enregistrer et passer au dossier ›")
        # Les pièces se déposent dès maintenant, et la fin s'enregistre à tout moment.
        self.assertContains(page, '<a href="?etape=3#etape-3">dès maintenant</a>', html=True)
        self.assertContains(page, "Mettre fin au mandat dès maintenant")
        page, modes = self.etapes(mandat, etape=3)
        self.assertEqual((modes[2], modes[3]), ("grisee", "ouverte"))
        self.assertContains(page, "Fermer le dossier")
        # ‹ ou le crayon : la rédaction se rouvre, › ramène à la signature.
        page, modes = self.etapes(mandat, etape=1)
        self.assertEqual((modes[1], modes[2]), ("ouverte", "grisee"))
        self.assertEqual(self.fleches(page), (None, 2))
        self.assertContains(page, "Enregistrer et passer à la signature ›")

    def test_mandat_signe_dossier_incomplet(self):
        """Signé, la rédaction se verrouille ; la signature se retire tant
        que le mandat n'est pas en gestion et qu'aucun bail n'est signé."""
        mandat = self.creer_mandat(date_signature=self.signe_le)
        page, modes = self.etapes(mandat, etape=1)
        self.assertEqual(modes, {1: "grisee", 2: "grisee", 3: "ouverte", 4: "a-venir", 5: "a-venir"})
        self.assertContains(page, "verrouillé, mandat signé")
        self.assertEqual((self.crayons(page), self.fleches(page)), ([2], (2, None)))
        reponse = self.client.post(mandat.get_absolute_url(), {"numero": "999"}, follow=True)
        self.assertContains(reponse, "Ce mandat est signé : pour le corriger, revenir d&#x27;abord à la signature.")
        page, modes = self.etapes(mandat, etape=2)
        self.assertEqual(modes[2], "ouverte")
        self.assertContains(page, "Retirer la signature…")
        reponse = self.client.post(reverse("mandats:retirer_signature", args=[mandat.pk]))
        self.assertRedirects(reponse, mandat.get_absolute_url() + "?etape=1#etape-1")
        mandat.refresh_from_db()
        self.assertFalse(mandat.signe)

    def test_en_gestion_puis_fin(self):
        mandat = self.creer_mandat(date_signature=self.signe_le)
        complet = mock.patch.object(Mandat, "pieces_manquantes", new_callable=mock.PropertyMock, return_value=[])
        with complet:
            page, modes = self.etapes(mandat)
            self.assertEqual(modes, {1: "grisee", 2: "grisee", 3: "grisee", 4: "ouverte", 5: "a-venir"})
            self.assertContains(page, "Mettre fin au mandat ›")
            # La signature ne se retire plus ; le dossier se rouvre, pour remplacer une pièce.
            self.assertEqual((self.crayons(page), self.fleches(page)), ([3], (3, None)))
            self.assertContains(page, 'class="verrou"', count=2)
            page, modes = self.etapes(mandat, etape=5)
            self.assertEqual(modes[5], "ouverte")
            fin = reverse("mandats:fin", args=[mandat.pk])
            aujourd_hui = timezone.localdate().isoformat()
            self.assertRedirects(self.client.post(fin, {"date_fin": aujourd_hui}), mandat.get_absolute_url() + "#etape-5")
            mandat.refresh_from_db()
            # Terminé : tout est verrouillé, seul un administrateur corrige.
            page, modes = self.etapes(mandat, etape=5)
            self.assertEqual(set(modes.values()), {"grisee"})
            self.assertEqual((self.crayons(page), self.fleches(page)), ([], (None, None)))
            self.assertNotContains(page, ">Corriger</a>")
            reponse = self.client.post(fin, {"date_fin": aujourd_hui}, follow=True)
            self.assertContains(reponse, "seul un administrateur peut corriger sa date de fin")
            admin = User.objects.create_user("admin")
            attribuer_role(admin, ADMINISTRATEURS)
            self.client.force_login(admin)
            self.assertContains(self.client.get(mandat.get_absolute_url()), ">Corriger</a>", count=2)
            page, modes = self.etapes(mandat, etape=5)
            self.assertEqual(modes[5], "ouverte")
            self.assertContains(page, "Enregistrer la correction")

    def test_bien_plus_gere(self):
        """Un bien décoché « Actif » (plus géré par l'agence) reste affiché
        dans la rédaction de son mandat, qui s'enregistre toujours."""
        mandat = self.creer_mandat()
        Propriete.objects.filter(pk=mandat.propriete_id).update(actif=False)
        page = self.client.get(mandat.get_absolute_url(), {"etape": 1})
        self.assertIn(f'<option value="{mandat.propriete_id}" selected>', str(page.context["form"]["propriete"]))
        form = page.context["form"]
        donnees = {nom: form[nom].value() for nom in form.fields if not form.fields[nom].disabled}
        donnees = {nom: "" if valeur is None else valeur for nom, valeur in donnees.items() if valeur is not False}
        donnees["mandant"] = "Monsieur DURAND Paul"
        self.assertRedirects(self.client.post(mandat.get_absolute_url(), donnees), mandat.get_absolute_url() + "#etape-2")


class ConformiteTests(Donnees, TestCase):
    """Loi Hoguet et code de la consommation : ce qui manque bloque la
    préparation du mandat, et le particulier qui signe hors de l'agence est
    informé de son droit de rétractation."""

    def setUp(self):
        self.client.force_login(creer_gestionnaire("gestion"))

    def preparer(self, mandat):
        return self.client.post(reverse("documents:preparer", args=["mandat", mandat.pk]), follow=True)

    def test_mentions_de_l_agence_manquantes(self):
        mandat = self.creer_mandat()
        remplir_agence(mediateur="", rcp="")
        reponse = self.preparer(mandat)
        self.assertContains(reponse, "Mentions légales de l&#x27;agence à compléter")
        self.assertContains(reponse, "médiateur de la consommation")
        self.assertFalse(mandat.versions_document.exists())
        self.assertEqual(self.client.get(reverse("mandats:word", args=[mandat.pk])).status_code, 302)

    def test_assurance_loyers_impayes_sans_orias(self):
        mandat = self.creer_mandat(assurance_loyers_impayes=True)
        remplir_agence(orias="")
        self.assertContains(self.preparer(mandat), "numéro ORIAS")
        mandat.assurance_loyers_impayes = False
        mandat.save()
        self.preparer(mandat)
        self.assertTrue(mandat.versions_document.exists())

    def test_lieu_de_signature_a_indiquer(self):
        mandat = self.creer_mandat(conclusion="")
        self.assertContains(self.preparer(mandat), "Indiquer qui est le mandant et où il signe")
        # Une société est un mandant professionnel : pas de rétractation.
        societe = Bailleur.objects.create(type=Bailleur.Type.PERSONNE_MORALE, nom="Foncière du Canal")
        propriete = self.creer_propriete(adresse="4 rue du Taur")
        propriete.bailleur = societe
        propriete.save()
        self.assertEqual(Mandat.depuis_propriete(propriete).conclusion, Mandat.Conclusion.PROFESSIONNEL)

    def test_retractation_hors_agence(self):
        en_agence = texte_docx(self.client.get(reverse("mandats:word", args=[self.creer_mandat().pk])).content)
        self.assertNotIn("Droit de rétractation", en_agence)
        self.assertNotIn("FORMULAIRE DE RÉTRACTATION", en_agence)
        self.assertIn("médiateur de la consommation", en_agence)
        self.assertIn("article 78 du décret", en_agence)
        mandat = self.creer_mandat(self.creer_propriete(adresse="4 rue du Taur"),
                                   conclusion=Mandat.Conclusion.HORS_ETABLISSEMENT, lieu_signature="Frouzins")
        texte = texte_docx(self.client.get(reverse("mandats:word", args=[mandat.pk])).content)
        self.assertIn("Droit de rétractation", texte)
        self.assertIn("quatorze jours à compter de la signature", texte)
        self.assertIn("FORMULAIRE DE RÉTRACTATION", texte)
        self.assertIn("Fait à Frouzins", texte)

    def test_modele_sans_les_mentions_obligatoires(self):
        from documents.views import verifier_modele

        document = Document(generation.MODELES_PAR_DEFAUT / "mandat_gestion.docx")
        # Ancien modèle de l'agence : les blocs de rétractation n'existent pas.
        dans_bloc = False
        for paragraphe in document.paragraphs:
            dans_bloc = dans_bloc or "if retractation" in paragraphe.text or "if formulaire_retractation" in paragraphe.text
            if dans_bloc:
                fin = "endif" in paragraphe.text
                paragraphe.text = ""
                dans_bloc = not fin
        flux = io.BytesIO()
        document.save(flux)
        message = verifier_modele("mandat_gestion.docx", flux.getvalue())
        self.assertIn("le droit de rétractation du mandant particulier", message)
        self.assertEqual(generation.defauts_modele("mandat_gestion.docx"), [])
