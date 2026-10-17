"""Import des contacts depuis un fichier Excel (.xlsx) ou CSV.

Une ligne = une personne ou une entreprise. Une personne déjà au carnet
(même nom et même prénom, sauf dates de naissance différentes) est complétée
au lieu d'être créée en double ; une cellule vide ne remplace jamais une
valeur déjà saisie. Chaque ligne passe par le formulaire de saisie d'un
contact : mêmes contrôles qu'à la main.
"""

import io
import re
from dataclasses import dataclass, field

from django.db import transaction
from django.forms.models import model_to_dict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from proprietes.tableur import AnnulerSimulation, Colonne, ErreurCellule, choix_de, convertir, lire_lignes, normaliser

from .forms import ContactForm
from .models import Contact, cle_identite


def _colonne(entete, champ, type="texte", obligatoire=False, largeur=18):
    return Colonne(entete, champ, type, modele=Contact, obligatoire=obligatoire, largeur=largeur)


COLONNES = [
    _colonne("Catégorie", "categorie", "choix", largeur=24),
    _colonne("Nature", "nature", "choix", largeur=12),
    _colonne("Civilité", "civilite", "choix", largeur=10),
    _colonne("Nom ou raison sociale", "nom", obligatoire=True, largeur=28),
    _colonne("Prénom", "prenom"),
    _colonne("Date de naissance", "date_naissance", "date", largeur=14),
    _colonne("Lieu de naissance", "lieu_naissance"),
    _colonne("E-mail", "email", largeur=26),
    _colonne("Téléphone", "telephone", largeur=16),
    _colonne("Adresse", "adresse", largeur=30),
    _colonne("Code postal", "code_postal", largeur=12),
    _colonne("Ville", "ville"),
    _colonne("Notes", "notes", largeur=30),
]

# Intitulés courants dans les fichiers tenus à la main ou exportés d'une messagerie.
AUTRES_ENTETES = {
    "nom": "Nom ou raison sociale", "raison sociale": "Nom ou raison sociale",
    "email": "E-mail", "mail": "E-mail", "courriel": "E-mail", "adresse e mail": "E-mail",
    "tel": "Téléphone", "telephone portable": "Téléphone", "portable": "Téléphone", "mobile": "Téléphone",
    "cp": "Code postal", "commune": "Ville", "civ": "Civilité", "type": "Catégorie",
}


# Façons courantes d'écrire une catégorie.
CATEGORIES_COURANTES = {
    "proprietaire": Contact.Categorie.BAILLEUR, "proprietaires": Contact.Categorie.BAILLEUR,
    "bailleurs": Contact.Categorie.BAILLEUR, "locataires": Contact.Categorie.LOCATAIRE,
    "garants": Contact.Categorie.GARANT, "caution": Contact.Categorie.GARANT,
}


@dataclass
class Rapport:
    crees: int = 0
    completes: int = 0
    erreurs: list = field(default_factory=list)  # (n° de ligne, message)
    colonnes_ignorees: list = field(default_factory=list)

    @property
    def total(self):
        return self.crees + self.completes


def importer(nom_fichier, contenu, simulation=False):
    """Importe le fichier. En simulation, tout est vérifié puis annulé."""
    rapport = Rapport()
    lignes = lire_lignes(nom_fichier, contenu)
    if not lignes:
        rapport.erreurs.append((1, "Le fichier est vide."))
        return rapport

    par_entete = {normaliser(c.entete): c for c in COLONNES}
    for autre, entete in AUTRES_ENTETES.items():
        par_entete.setdefault(autre, par_entete[normaliser(entete)])
    index = {}
    for position, entete in enumerate(lignes[0]):
        if entete is None or not str(entete).strip():
            continue
        colonne = par_entete.get(normaliser(entete))
        if colonne and colonne not in index:
            index[colonne] = position
        else:
            rapport.colonnes_ignorees.append(str(entete))
    manquantes = [c.entete for c in COLONNES if c.obligatoire and c not in index]
    if manquantes:
        rapport.erreurs.append((1, "Colonnes obligatoires absentes : " + ", ".join(manquantes)))
        return rapport

    try:
        with transaction.atomic():
            for numero, cellules in enumerate(lignes[1:], start=2):
                if all(c is None or str(c).strip() == "" for c in cellules):
                    continue
                try:
                    with transaction.atomic():
                        _importer_ligne(cellules, index, rapport)
                except ErreurCellule as erreur:
                    rapport.erreurs.append((numero, str(erreur)))
            if simulation:
                raise AnnulerSimulation
    except AnnulerSimulation:
        pass
    return rapport


def _importer_ligne(cellules, index, rapport):
    valeurs = {}
    for colonne, position in index.items():
        brut = cellules[position] if position < len(cellules) else None
        if colonne.champ == "categorie" and normaliser(brut or "") in CATEGORIES_COURANTES:
            brut = CATEGORIES_COURANTES[normaliser(brut)]
        try:
            valeur = convertir(colonne, brut)
        except ErreurCellule as erreur:
            raise ErreurCellule(f"{colonne.entete} : {erreur}") from None
        if valeur is not None:
            valeurs[colonne.champ] = valeur
    # Excel lit « 0612345678 » et « 01000 » comme des nombres, sans le zéro de tête.
    if re.fullmatch(r"[1-9]\d{8}", valeurs.get("telephone", "")):
        valeurs["telephone"] = "0" + valeurs["telephone"]
    if re.fullmatch(r"\d{4}", valeurs.get("code_postal", "")):
        valeurs["code_postal"] = "0" + valeurs["code_postal"]
    if not valeurs.get("nom"):
        raise ErreurCellule("Nom ou raison sociale : valeur obligatoire")
    if valeurs.get("nature") is None and not valeurs.get("prenom") and not valeurs.get("civilite"):
        # Sans prénom ni civilité, c'est une entreprise (SCI, syndic, artisan…).
        valeurs["nature"] = Contact.Nature.ENTREPRISE

    contact = _trouver(valeurs)
    donnees = model_to_dict(contact, fields=ContactForm.Meta.fields) if contact else {}
    donnees = {champ: valeur for champ, valeur in donnees.items() if valeur not in (None, "")}
    if contact and contact.categorie and "categorie" in valeurs:
        valeurs.pop("categorie")  # la catégorie indiquée à la création ne change pas
    donnees.update(valeurs)
    donnees.setdefault("nature", Contact.Nature.PERSONNE)

    form = ContactForm(donnees, instance=contact)
    form.fields["categorie"].error_messages["required"] = "à indiquer pour un contact qui n'est pas encore au carnet."
    if not form.is_valid():
        raise ErreurCellule(_message(form))
    objet = form.save()
    if contact and contact.automatique:
        # Une personne inscrite au carnet par l'agence ne s'efface plus avec sa candidature.
        objet.automatique = False
        objet.save(update_fields=["automatique"])
    objet.reporter_sur_bailleurs()
    if contact:
        rapport.completes += 1
    else:
        rapport.crees += 1


def _trouver(valeurs):
    """Contact déjà au carnet : même nom et même prénom, sauf dates de naissance différentes."""
    naissance = valeurs.get("date_naissance")
    memes = Contact.objects.filter(cle=cle_identite(valeurs["nom"], valeurs.get("prenom", "")))
    return next((c for c in memes if not (c.date_naissance and naissance and c.date_naissance != naissance)), None)


def _message(form):
    entetes = {c.champ: c.entete for c in COLONNES}
    return " ; ".join(
        f"{entetes.get(champ, '')} : {' '.join(erreurs)}" if champ in entetes else " ".join(erreurs)
        for champ, erreurs in form.errors.items()
    )


def classeur():
    """Modèle vide à remplir."""
    wb = Workbook()
    feuille = wb.active
    feuille.title = "Contacts"
    feuille.append([c.entete for c in COLONNES])
    for position, colonne in enumerate(COLONNES, start=1):
        cellule = feuille.cell(row=1, column=position)
        cellule.font = Font(bold=True)
        if colonne.obligatoire or colonne.champ == "categorie":
            cellule.fill = PatternFill("solid", fgColor="FDE9C9")
        feuille.column_dimensions[get_column_letter(position)].width = colonne.largeur
        if colonne.type == "choix":
            libelles = [libelle for _, libelle in choix_de(colonne)]
            validation = DataValidation(type="list", formula1='"' + ",".join(libelles) + '"', allow_blank=True)
            lettre = get_column_letter(position)
            validation.add(f"{lettre}2:{lettre}2000")
            feuille.add_data_validation(validation)
    feuille.freeze_panes = "A2"

    aide = wb.create_sheet("Mode d'emploi")
    aide.column_dimensions["A"].width = 22
    aide.column_dimensions["B"].width = 90
    for ligne in (
        ("Une ligne", "Une personne ou une entreprise."),
        ("Obligatoire", "Le nom, et la catégorie pour un contact qui n'est pas encore au carnet."),
        ("Catégorie", ", ".join(libelle for _, libelle in Contact.Categorie.choices)
         + ". Un propriétaire reçoit aussi sa fiche de bailleur."),
        ("Entreprise", "Sans prénom ni civilité, la ligne est une entreprise (SCI, société…)."),
        ("Déjà au carnet", "Un contact de même nom et prénom est complété ; une cellule vide n'efface rien."),
    ):
        aide.append(ligne)
        aide.cell(row=aide.max_row, column=1).font = Font(bold=True)

    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
