import io
import shutil
import tempfile
import urllib.error
from pathlib import Path
from unittest import mock

from core.testing import creer_gestionnaire

from django.conf import settings
from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document
from jinja2.exceptions import SecurityError

from comptes.roles import ADMINISTRATEURS, attribuer_role
from mandats.tests import Donnees, texte_docx

from . import google
from .generation import MODELES_PAR_DEFAUT, generer_docx
from .models import Modele

LIEN = "https://docs.google.com/document/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit?usp=sharing"


def modele_mandat(*paragraphes):
    """Modèle de mandat d'origine augmenté de quelques paragraphes."""
    document = Document(MODELES_PAR_DEFAUT / "mandat_gestion.docx")
    for texte in paragraphes:
        document.add_paragraph(texte)
    flux = io.BytesIO()
    document.save(flux)
    return flux.getvalue()


class ReponseHttp(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()


class GoogleTests(TestCase):
    def test_identifiant(self):
        self.assertEqual(google.identifiant(LIEN), "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789")
        self.assertEqual(google.identifiant("https://docs.google.com/document/u/0/d/1AbCdEfGhIjKlMnOpQrStUvWx/edit"),
                         "1AbCdEfGhIjKlMnOpQrStUvWx")
        for lien in ["https://drive.google.com/file/d/1AbCdEfGhIjKlMnOpQrStUvWx/view",
                     "https://docs.google.com.exemple.fr/document/d/1AbCdEfGhIjKlMnOpQrStUvWx", "", None]:
            with self.subTest(lien=lien):
                self.assertIsNone(google.identifiant(lien))

    def test_telechargement(self):
        with mock.patch("urllib.request.urlopen", return_value=ReponseHttp(b"PK\x03\x04contenu")) as appel:
            self.assertEqual(google.telecharger(LIEN), b"PK\x03\x04contenu")
        self.assertEqual(appel.call_args.args[0].full_url,
                         "https://docs.google.com/document/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/export?format=docx")
        # Document non partagé : Google renvoie sa page de connexion.
        with mock.patch("urllib.request.urlopen", return_value=ReponseHttp(b"<!doctype html>")), \
                self.assertRaisesMessage(google.LectureImpossible, "n'est pas partagé"):
            google.telecharger(LIEN)
        erreur = urllib.error.HTTPError(LIEN, 404, "Not Found", {}, None)
        with mock.patch("urllib.request.urlopen", side_effect=erreur), \
                self.assertRaisesMessage(google.LectureImpossible, "Google refuse l'accès"):
            google.telecharger(LIEN)
        with mock.patch("urllib.request.urlopen", side_effect=urllib.error.URLError("hors ligne")), \
                self.assertRaisesMessage(google.LectureImpossible, "injoignable"):
            google.telecharger(LIEN)


@override_settings(MFA_OBLIGATOIRE=False)
class ModelesTests(Donnees, TestCase):
    def setUp(self):
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier)
        reglage = override_settings(MEDIA_ROOT=dossier)
        reglage.enable()
        self.addCleanup(reglage.disable)
        admin = User.objects.create_user("elodie", first_name="Elodie")
        attribuer_role(admin, ADMINISTRATEURS)
        self.client.force_login(admin)
        self.mandat = self.creer_mandat()
        self.url = reverse("documents:reprendre", args=["mandat_gestion.docx"])

    def reprendre(self, contenu, lien=LIEN):
        with mock.patch("documents.google.telecharger", return_value=contenu):
            return self.client.post(self.url, {"lien_google": lien}, follow=True)

    def mettre_en_service(self, **donnees):
        return self.client.post(reverse("documents:mettre_en_service", args=["mandat_gestion.docx"]), donnees, follow=True)

    def ouvrir_essai(self):
        with mock.patch("documents.views.convertir_pdf", return_value=b"%PDF-1.7"):
            return self.client.get(reverse("documents:essai_version", args=["mandat_gestion.docx"]))

    def texte_mandat(self):
        return texte_docx(self.client.get(reverse("mandats:word", args=[self.mandat.pk])).content)

    def test_reprise_depuis_google_docs(self):
        reponse = self.reprendre(modele_mandat("Clause ajoutée dans Google Docs pour {{ mandant }}."))
        self.assertContains(reponse, "repris de Google Docs. Ouvrez le document d&#x27;essai")
        self.assertContains(reponse, "Essai prêt (repris de Google Docs)")
        # Rien ne change tant que l'essai n'est pas ouvert puis mis en service.
        self.assertNotIn("Clause ajoutée", self.texte_mandat())
        self.assertContains(self.mettre_en_service(relu="1", reference="la loi Hoguet"), "Ouvrir d&#x27;abord le document d&#x27;essai")
        self.assertNotIn("Clause ajoutée", self.texte_mandat())
        self.assertEqual(self.ouvrir_essai().content, b"%PDF-1.7")
        reponse = self.mettre_en_service(relu="1", reference="la loi Hoguet")
        self.assertContains(reponse, "mis en service, relecture notée")
        self.assertContains(reponse, "Repris de Google Docs, le")
        self.assertContains(reponse, f'value="{LIEN}"')
        self.assertIn("Clause ajoutée dans Google Docs pour Monsieur DURAND Paul", self.texte_mandat())
        modele = Modele.objects.get(nom="mandat_gestion.docx")
        self.assertEqual((modele.origine, modele.lien_google), (Modele.Origine.GOOGLE, LIEN))
        self.assertEqual((modele.relu_par.username, modele.relu_reference), ("elodie", "la loi Hoguet"))
        # Le lien est gardé : la version suivante se reprend sans le recoller,
        # et sans la case cochée la relecture reste à faire.
        with mock.patch("documents.google.telecharger", return_value=modele_mandat("Deuxième version.")):
            self.client.post(self.url)
        self.ouvrir_essai()
        self.assertContains(self.mettre_en_service(reference="la loi Hoguet"), "mis en service, à relire")
        self.assertIn("Deuxième version.", self.texte_mandat())
        # Un essai abandonné ne change rien.
        self.reprendre(modele_mandat("Troisième version."))
        self.client.post(reverse("documents:abandonner", args=["mandat_gestion.docx"]))
        self.ouvrir_essai()
        self.mettre_en_service()
        self.assertIn("Deuxième version.", self.texte_mandat())
        # Retour à la version précédente, puis au modèle d'origine.
        self.client.post(reverse("documents:precedent", args=["mandat_gestion.docx"]))
        self.assertIn("Clause ajoutée dans Google Docs", self.texte_mandat())
        self.client.post(reverse("documents:retablir", args=["mandat_gestion.docx"]))
        self.assertNotIn("Clause ajoutée", self.texte_mandat())

    def test_modele_refuse(self):
        for paragraphe, message in [
            ("{% if mandant %}sans fin", "Une balise est mal écrite"),
            ("{{ balise_inconnue }}", "Balises inconnues : balise_inconnue"),
            ("{{ mandant.upper(1, 2, 3) }}", "Le remplissage d&#x27;essai échoue"),
        ]:
            with self.subTest(paragraphe=paragraphe):
                self.assertContains(self.reprendre(modele_mandat(paragraphe)), message)
                self.assertNotIn("sans fin", self.texte_mandat())
        self.assertFalse(Modele.objects.filter(origine=Modele.Origine.GOOGLE).exists())

    def test_modele_ne_peut_pas_executer_de_code(self):
        # Une balise qui remonte aux classes de Python passerait le contrôle des
        # balises inconnues : c'est l'environnement protégé qui la bloque.
        piege = "{{ ''.__class__.__mro__[1].__subclasses__()|length }}"
        reponse = self.reprendre(modele_mandat(piege))
        self.assertContains(reponse, "Le remplissage d&#x27;essai échoue")
        self.assertFalse(Modele.objects.filter(origine=Modele.Origine.GOOGLE).exists())
        # Même déposé directement sur le disque, le modèle ne produit rien.
        depose = Path(settings.MEDIA_ROOT) / "modeles" / "mandat_gestion.docx"
        depose.parent.mkdir(parents=True, exist_ok=True)
        depose.write_bytes(modele_mandat(piege))
        with self.assertRaises(SecurityError):
            generer_docx("mandat_gestion.docx", self.mandat.contexte_document())

    def test_lien_ou_partage_incorrect(self):
        reponse = self.client.post(self.url, {"lien_google": "https://exemple.fr/doc"}, follow=True)
        self.assertContains(reponse, "Coller le lien du Google Doc")
        with mock.patch("documents.google.telecharger", side_effect=google.LectureImpossible("Pas partagé.")):
            reponse = self.client.post(self.url, {"lien_google": LIEN}, follow=True)
        self.assertContains(reponse, "n&#x27;a pas été repris. Pas partagé.")
        # Le lien est retenu pour réessayer une fois le partage corrigé.
        self.assertEqual(Modele.objects.get(nom="mandat_gestion.docx").lien_google, LIEN)

    def test_document_essai(self):
        with mock.patch("documents.views.convertir_pdf", return_value=b"%PDF-1.7"):
            reponse = self.client.get(reverse("documents:essai", args=["mandat_gestion.docx"]))
        self.assertEqual(reponse.content, b"%PDF-1.7")
        reponse = self.client.get(reverse("documents:essai", args=["bail_habitation.docx"]), follow=True)
        self.assertContains(reponse, "Aucun bail d&#x27;habitation (nu ou meublé) enregistré")

    def test_reserve_aux_administrateurs(self):
        self.client.force_login(creer_gestionnaire("gestion"))
        self.assertEqual(self.reprendre(modele_mandat()).status_code, 403)
        self.assertEqual(self.client.get(reverse("documents:essai", args=["mandat_gestion.docx"])).status_code, 403)


class ReglesTests(TestCase):
    def setUp(self):
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier)
        reglage = override_settings(MEDIA_ROOT=dossier)
        reglage.enable()
        self.addCleanup(reglage.disable)
        self.admin = User.objects.create_user("admin", first_name="Élodie")
        attribuer_role(self.admin, ADMINISTRATEURS)
        self.client.force_login(self.admin)

    def test_modeles_d_origine_respectent_toutes_les_regles(self):
        from .regles import A_CORRIGER, registre
        lignes = [ligne for groupe in registre().values() for ligne in groupe]
        self.assertEqual([ligne["regle"].libelle for ligne in lignes if ligne["etat"] == A_CORRIGER], [])

    def test_modele_depose_avec_clause_perimee_est_a_corriger(self):
        from .regles import A_CORRIGER, registre
        depose = Path(settings.MEDIA_ROOT) / "modeles" / "bail_habitation.docx"
        depose.parent.mkdir(parents=True)
        document = Document(MODELES_PAR_DEFAUT / "bail_habitation.docx")
        document.add_paragraph("Résiliation de plein droit deux mois après un commandement de payer.")
        document.save(depose)
        etats = {ligne["regle"].libelle: ligne["etat"] for ligne in registre()["Bail"]}
        self.assertEqual(etats["Clause résolutoire"], A_CORRIGER)
        reponse = self.client.get(reverse("administration"))
        self.assertContains(reponse, "1 à corriger")

    def test_regle_verifiee_il_y_a_plus_d_un_an_est_a_reverifier(self):
        from datetime import timedelta
        from .regles import A_REVERIFIER, REGLES, etat
        regle = REGLES[1]
        self.assertEqual(etat(regle, {}, regle.verifiee_le + timedelta(days=366))[0], A_REVERIFIER)

    def test_page_reservee_aux_administrateurs(self):
        self.assertContains(self.client.get(reverse("documents:regles")), "Dépôt de garantie maximal")
        self.client.force_login(creer_gestionnaire("gestion"))
        self.assertEqual(self.client.get(reverse("documents:regles")).status_code, 403)

    def test_marquer_un_modele_relu(self):
        url = reverse("documents:relu", args=["mandat_gestion.docx"])
        self.client.post(url, {"reference": "  "})
        self.assertFalse(Modele.objects.filter(relu_le__isnull=False).exists())
        self.client.post(url, {"reference": "la loi Hoguet et le décret 72-678"})
        modele = Modele.objects.get(nom="mandat_gestion.docx")
        self.assertEqual((modele.relu_par, modele.relu_reference), (self.admin, "la loi Hoguet et le décret 72-678"))
        page = self.client.get(reverse("documents:modele", args=["mandat_gestion.docx"]))
        self.assertContains(page, "d'après la loi Hoguet et le décret 72-678")
        self.assertContains(self.client.get(reverse("documents:modele", args=["bail_habitation.docx"])), "Pas encore relu")


class CorrectionTests(TestCase):
    def setUp(self):
        dossier = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, dossier)
        reglage = override_settings(MEDIA_ROOT=dossier)
        reglage.enable()
        self.addCleanup(reglage.disable)
        self.admin = User.objects.create_user("admin", first_name="Élodie")
        attribuer_role(self.admin, ADMINISTRATEURS)
        self.client.force_login(self.admin)
        # Modèle de l'agence déposé avant que la clause ne soit périmée, la
        # phrase répartie sur deux morceaux de mise en forme.
        document = Document(MODELES_PAR_DEFAUT / "bail_habitation.docx")
        paragraphe = document.add_paragraph("Résilié de plein droit deux ")
        paragraphe.add_run("mois").bold = True
        paragraphe.add_run(" après un commandement de payer resté sans effet.")
        self.depose = Path(settings.MEDIA_ROOT) / "modeles" / "bail_habitation.docx"
        self.depose.parent.mkdir(parents=True)
        document.save(self.depose)
        self.page = reverse("documents:corriger", args=["bail_habitation.docx"])
        self.appliquer = reverse("documents:appliquer_correction", args=["bail_habitation.docx"])

    def test_regle_propose_la_correction(self):
        self.assertContains(self.client.get(reverse("documents:regles")), self.page)
        page = self.client.get(self.page)
        self.assertContains(page, "<del>deux mois après un commandement de payer</del>", html=False)
        self.assertContains(page, "<ins>six semaines après un commandement de payer</ins>", html=False)
        self.assertContains(page, "disabled")

    def test_mise_en_service_exige_l_essai(self):
        self.client.post(self.appliquer)
        self.assertIn("deux mois après", texte_docx(self.depose.read_bytes()))
        self.client.get(reverse("documents:essai_correction", args=["bail_habitation.docx"]))
        self.assertNotContains(self.client.get(self.page), "disabled")
        self.client.post(self.appliquer)
        texte = texte_docx(self.depose.read_bytes())
        self.assertIn("Résilié de plein droit six semaines après un commandement de payer resté sans effet.", texte)
        self.assertNotIn("deux mois après", texte)
        modele = Modele.objects.get(nom="bail_habitation.docx")
        self.assertEqual(modele.relu_par, self.admin)
        self.assertIn("six semaines", modele.relu_reference)
        precedent = Path(settings.MEDIA_ROOT) / "modeles" / "precedent" / "bail_habitation.docx"
        self.assertIn("deux mois après", texte_docx(precedent.read_bytes()))

    def test_essai_perime_si_le_modele_change(self):
        self.client.get(reverse("documents:essai_correction", args=["bail_habitation.docx"]))
        document = Document(self.depose)
        document.add_paragraph("Ajout de l'agence.")
        document.save(self.depose)
        self.client.post(self.appliquer)
        self.assertIn("deux mois après", texte_docx(self.depose.read_bytes()))

    def test_sans_clause_perimee_rien_a_corriger(self):
        self.depose.unlink()
        self.assertRedirects(self.client.get(self.page), reverse("documents:modele", args=["bail_habitation.docx"]))

    def test_liste_et_tableau_de_bord_signalent_le_modele(self):
        fiche = reverse("documents:modele", args=["bail_habitation.docx"])
        liste = self.client.get(reverse("documents:modeles"))
        self.assertContains(liste, "Clause résolutoire : six semaines après un commandement de payer")
        self.assertContains(liste, ">Corriger</a>", html=False)
        self.assertContains(self.client.get(fiche), self.page)
        alerte = "ne respecte plus une règle légale"
        self.assertContains(self.client.get(reverse("accueil")), alerte)
        # Seul un administrateur peut corriger : l'alerte ne s'adresse qu'à lui.
        self.client.force_login(creer_gestionnaire("gestion"))
        self.assertNotContains(self.client.get(reverse("accueil")), alerte)

    def test_reserve_aux_administrateurs(self):
        self.client.force_login(creer_gestionnaire("gestion"))
        self.assertEqual(self.client.get(self.page).status_code, 403)
        self.assertEqual(self.client.get(reverse("documents:modele", args=["bail_habitation.docx"])).status_code, 403)
