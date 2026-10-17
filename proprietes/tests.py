import io
from decimal import Decimal

from core.testing import creer_gestionnaire

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from . import tableur
import datetime as dt
import tempfile

from .models import Bailleur, Immeuble, Piece, Propriete

ENTETES = [
    "N° de lot", "Réf. ICS bailleur", "Type de bailleur", "Bailleur (nom ou raison sociale)",
    "Bailleur prénom", "Usage", "Adresse", "Complément", "Code postal", "Ville",
    "Surface (m²)", "Pièces principales", "Classe DPE", "Zone tendue", "Date DPE",
]


def ligne(ref="L001", **autres):
    valeurs = {
        "N° de lot": ref, "Réf. ICS bailleur": "B01", "Type de bailleur": "Personne physique",
        "Bailleur (nom ou raison sociale)": "Durand", "Bailleur prénom": "Anne",
        "Usage": "Habitation", "Adresse": "12 rue des Lois", "Complément": "Bât. A, 2e étage",
        "Code postal": "31000", "Ville": "Toulouse", "Surface (m²)": 45.5,
        "Pièces principales": 2, "Classe DPE": "D", "Zone tendue": "oui", "Date DPE": "15/03/2024",
    }
    valeurs.update(autres)
    return [valeurs.get(entete) for entete in ENTETES]


def xlsx(*lignes, entetes=ENTETES):
    wb = Workbook()
    wb.active.append(entetes)
    for valeurs in lignes:
        wb.active.append(valeurs)
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


class ImportTests(TestCase):
    def test_import_cree_propriete_et_bailleur(self):
        rapport = tableur.importer("proprietes.xlsx", xlsx(ligne()))
        self.assertEqual((rapport.crees, rapport.bailleurs_crees, rapport.erreurs), (1, 1, []))
        propriete = Propriete.objects.get(numero_lot="L001")
        self.assertEqual(propriete.surface, Decimal("45.50"))
        self.assertEqual(propriete.classe_dpe, "D")
        self.assertTrue(propriete.zone_tendue)
        self.assertEqual(str(propriete.date_dpe), "2024-03-15")
        self.assertEqual(propriete.bailleur.ref_ics, "B01")

    def test_import_sans_reference_ics(self):
        """La référence ICS est facultative : le lot se retrouve alors par
        son adresse, son complément et son code postal."""
        rapport = tableur.importer("proprietes.xlsx", xlsx(
            ligne(None), ligne(None, **{"Complément": "Bât. B", "Réf. ICS bailleur": None})))
        self.assertEqual((rapport.erreurs, rapport.crees), ([], 2))
        self.assertEqual(Propriete.objects.filter(numero_lot__isnull=True).count(), 2)
        rapport = tableur.importer("proprietes.xlsx", xlsx(ligne(None, **{"Classe DPE": "C"})))
        self.assertEqual((rapport.crees, rapport.mis_a_jour), (0, 1))
        self.assertEqual(Propriete.objects.get(complement="Bât. A, 2e étage").classe_dpe, "C")
        # La référence peut être ajoutée plus tard sur le même lot.
        rapport = tableur.importer("proprietes.xlsx", xlsx(ligne("L001")))
        self.assertEqual((rapport.crees, rapport.mis_a_jour), (0, 1))
        self.assertEqual(Propriete.objects.get(numero_lot="L001").complement, "Bât. A, 2e étage")

    def test_ancienne_colonne_ref_ics_lot(self):
        """Un fichier préparé avant le renommage garde sa colonne « Réf. ICS lot »."""
        anciennes = ["Réf. ICS lot" if entete == "N° de lot" else entete for entete in ENTETES]
        tableur.importer("proprietes.xlsx", xlsx(ligne("L001"), entetes=anciennes))
        self.assertTrue(Propriete.objects.filter(numero_lot="L001").exists())

    def test_reimport_met_a_jour_sans_effacer_les_cellules_vides(self):
        tableur.importer("proprietes.xlsx", xlsx(ligne()))
        rapport = tableur.importer("proprietes.xlsx", xlsx(ligne(**{"Surface (m²)": None, "Classe DPE": "C"})))
        self.assertEqual((rapport.crees, rapport.mis_a_jour, rapport.bailleurs_crees), (0, 1, 0))
        propriete = Propriete.objects.get(numero_lot="L001")
        self.assertEqual(propriete.surface, Decimal("45.50"))
        self.assertEqual(propriete.classe_dpe, "C")
        self.assertEqual(Bailleur.objects.count(), 1)

    def test_erreurs_signalees_par_ligne_sans_bloquer_les_autres(self):
        rapport = tableur.importer("proprietes.xlsx", xlsx(
            ligne("L001"),
            ligne("L002", **{"Classe DPE": "Z"}),
            ligne("L003", **{"Surface (m²)": "quarante"}),
            ligne("L004", **{"Ville": None}),
        ))
        self.assertEqual(rapport.crees, 1)
        self.assertEqual([numero for numero, _ in rapport.erreurs], [3, 4, 5])
        self.assertIn("Classe DPE", rapport.erreurs[0][1])
        self.assertIn("nombre", rapport.erreurs[1][1])
        self.assertIn("obligatoire", rapport.erreurs[2][1])
        self.assertEqual(list(Propriete.objects.values_list("numero_lot", flat=True)), ["L001"])

    def test_simulation_n_enregistre_rien(self):
        rapport = tableur.importer("proprietes.xlsx", xlsx(ligne()), simulation=True)
        self.assertEqual(rapport.crees, 1)
        self.assertFalse(Propriete.objects.exists())
        self.assertFalse(Bailleur.objects.exists())

    def test_colonne_obligatoire_absente(self):
        rapport = tableur.importer("proprietes.xlsx", xlsx(["X"], entetes=["Adresse"]))
        self.assertIn("Code postal", rapport.erreurs[0][1])
        self.assertNotIn("N° de lot", rapport.erreurs[0][1])

    def test_csv_excel_francais(self):
        contenu = (
            "Ref ICS lot;Bailleur (nom ou raison sociale);Adresse;Code postal;Ville;Surface (m²);Usage;Meublé\n"
            "L010;Société Générale Immo;5 allée Jean Jaurès;31000;Toulouse;32,5;Local professionnel;non\n"
        ).encode("cp1252")
        rapport = tableur.importer("export.csv", contenu)
        self.assertEqual(rapport.erreurs, [])
        propriete = Propriete.objects.get(numero_lot="L010")
        self.assertEqual(propriete.surface, Decimal("32.50"))
        self.assertEqual(propriete.usage, Propriete.Usage.PROFESSIONNEL)
        self.assertEqual(propriete.bailleur.nom, "Société Générale Immo")

    def test_bailleur_retrouve_par_nom_sans_reference(self):
        Bailleur.objects.create(nom="Martin", prenom="Paul")
        tableur.importer("proprietes.xlsx", xlsx(ligne(**{
            "Réf. ICS bailleur": None, "Bailleur (nom ou raison sociale)": "MARTIN", "Bailleur prénom": "paul",
        })))
        self.assertEqual(Bailleur.objects.count(), 1)

    def test_export_reimportable(self):
        tableur.importer("proprietes.xlsx", xlsx(ligne("L001"), ligne("L002")))
        contenu = tableur.classeur(Propriete.objects.select_related("bailleur"))
        feuille = load_workbook(io.BytesIO(contenu)).active
        self.assertEqual(feuille.max_row, 3)
        Propriete.objects.update(surface=None)
        rapport = tableur.importer("export.xlsx", contenu)
        self.assertEqual((rapport.mis_a_jour, rapport.erreurs), (2, []))
        self.assertEqual(Propriete.objects.get(numero_lot="L001").surface, Decimal("45.50"))

    def test_modele_vide_a_listes_deroulantes(self):
        feuille = load_workbook(io.BytesIO(tableur.classeur())).active
        self.assertEqual(feuille.max_row, 1)
        self.assertTrue(feuille.data_validations.dataValidation)


class ProprieteTests(TestCase):
    def test_logement_g_signale(self):
        bailleur = Bailleur.objects.create(nom="Durand")
        propriete = Propriete(bailleur=bailleur, adresse="1 rue", code_postal="31000", ville="Toulouse", classe_dpe="G")
        self.assertTrue(propriete.location_interdite)
        propriete.usage = Propriete.Usage.COMMERCIAL
        self.assertFalse(propriete.location_interdite)

    def test_numero_de_lot_unique_dans_son_immeuble(self):
        """Le n° de lot est celui de l'immeuble : le même numéro se retrouve
        d'un immeuble à l'autre, jamais deux fois dans le même (#160)."""
        from django.core.exceptions import ValidationError

        bailleur = Bailleur.objects.create(nom="Durand")
        pins, canal = (Immeuble.objects.create(nom=nom, adresse="1 rue", code_postal="31000", ville="Toulouse")
                       for nom in ("Les Pins", "Le Canal"))
        valeurs = {"bailleur": bailleur, "adresse": "1 rue", "code_postal": "31000", "ville": "Toulouse", "numero_lot": "12"}
        Propriete.objects.create(immeuble=pins, **valeurs)
        Propriete(immeuble=canal, **valeurs).full_clean()
        with self.assertRaisesMessage(ValidationError, "Un lot de cet immeuble porte déjà ce numéro."):
            Propriete(immeuble=pins, **valeurs).full_clean()

    def test_duree_bail_selon_bailleur(self):
        self.assertEqual(Bailleur(type=Bailleur.Type.SCI_FAMILIALE).duree_bail_nu_ans, 3)
        self.assertEqual(Bailleur(type=Bailleur.Type.PERSONNE_MORALE).duree_bail_nu_ans, 6)


@override_settings(MFA_OBLIGATOIRE=False)
class VuesTests(TestCase):
    def setUp(self):
        self.client.force_login(creer_gestionnaire("gestion"))

    def test_parcours_import_verification_puis_confirmation(self):
        fichier = SimpleUploadedFile("proprietes.xlsx", xlsx(ligne("L001"), ligne("L002", **{"Classe DPE": "Z"})))
        reponse = self.client.post(reverse("proprietes:importer"), {"fichier": fichier})
        self.assertContains(reponse, "1 bien(s) à créer")
        self.assertContains(reponse, "Classe DPE")
        self.assertFalse(Propriete.objects.exists())

        reponse = self.client.post(reverse("proprietes:importer_confirmer"), follow=True)
        self.assertContains(reponse, "1 bien(s) créé(s)")
        self.assertTrue(Propriete.objects.filter(numero_lot="L001").exists())

        # Le fichier n'est importé qu'une fois.
        reponse = self.client.post(reverse("proprietes:importer_confirmer"), follow=True)
        self.assertContains(reponse, "Aucun fichier en attente")

    def test_fichier_illisible(self):
        fichier = SimpleUploadedFile("proprietes.xlsx", b"pas un classeur")
        reponse = self.client.post(reverse("proprietes:importer"), {"fichier": fichier})
        self.assertContains(reponse, "Impossible de lire ce fichier")

    def test_liste_recherche_et_export(self):
        tableur.importer("proprietes.xlsx", xlsx(ligne("L001"), ligne("L002", Ville="Blagnac", **{"Code postal": "31700"})))
        reponse = self.client.get(reverse("proprietes:liste"), {"q": "blagnac"})
        self.assertContains(reponse, "1 propriété(s)")
        reponse = self.client.get(reverse("proprietes:exporter"), {"q": "blagnac"})
        self.assertEqual(load_workbook(io.BytesIO(reponse.content)).active.max_row, 2)

    def test_creation_propriete_et_bailleur_sans_reference_ics(self):
        for nom in ("Durand", "Martin"):
            reponse = self.client.post(reverse("proprietes:bailleur_creer"), {"type": "physique", "nom": nom})
            self.assertEqual(reponse.status_code, 302)
        self.assertEqual(Bailleur.objects.filter(ref_ics__isnull=True).count(), 2)

        bailleur = Bailleur.objects.first()
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {
            "bailleur": bailleur.pk, "usage": "habitation", "actif": "on",
            "adresse": "3 place du Capitole", "code_postal": "31000", "ville": "Toulouse",
        })
        propriete = Propriete.objects.get()
        self.assertRedirects(reponse, propriete.get_absolute_url())
        reponse = self.client.get(propriete.get_absolute_url())
        self.assertContains(reponse, "À compléter avant de rédiger un bail")
        self.assertContains(reponse, "gestion")  # auteur dans l'historique

    def test_nouvelle_propriete_avec_nouveau_bailleur(self):
        """Un bailleur absent de la liste se crée sur la page de la propriété."""
        page = self.client.get(reverse("proprietes:propriete_creer"))
        self.assertContains(page, "Nouveau bailleur (ci-dessous)")
        self.assertContains(page, 'name="bailleur-nom"')
        propriete = {"usage": "habitation", "actif": "on", "adresse": "3 place du Capitole",
                     "code_postal": "31000", "ville": "Toulouse"}
        # Ni bailleur choisi ni bailleur saisi : rien n'est créé.
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {**propriete, "bailleur-type": "physique"})
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("nom", reponse.context["nouveau_bailleur"].errors)
        self.assertFalse(Propriete.objects.exists() or Bailleur.objects.exists())
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {
            **propriete, "bailleur-type": "physique", "bailleur-civilite": "Mme", "bailleur-nom": "Garcia",
            "bailleur-prenom": "Ines", "bailleur-ville": "Toulouse",
        })
        nouvelle = Propriete.objects.get()
        self.assertRedirects(reponse, nouvelle.get_absolute_url())
        self.assertEqual((nouvelle.bailleur.nom, nouvelle.bailleur.prenom, nouvelle.bailleur.ville), ("Garcia", "Ines", "Toulouse"))
        # Bailleur existant choisi : les champs du nouveau sont ignorés.
        self.client.post(reverse("proprietes:propriete_creer"), {
            **propriete, "adresse": "5 rue du Taur", "bailleur": nouvelle.bailleur.pk, "bailleur-nom": "Ignoré",
        })
        self.assertEqual((Bailleur.objects.count(), nouvelle.bailleur.proprietes.count()), (1, 2))
        # À la modification, le bailleur reste obligatoire et rien d'autre n'est proposé.
        page = self.client.get(reverse("proprietes:propriete_modifier", args=[nouvelle.pk]))
        self.assertNotContains(page, 'name="bailleur-nom"')
        reponse = self.client.post(reverse("proprietes:propriete_modifier", args=[nouvelle.pk]), propriete)
        self.assertIn("bailleur", reponse.context["form"].errors)

    def immeuble(self, **autres):
        """Saisie d'un immeuble de deux lots, la troisième ligne laissée vide."""
        return {
            "immeuble-nom": "Résidence du Canal", "immeuble-adresse": "8 quai de Brienne",
            "immeuble-code_postal": "31000", "immeuble-ville": "Toulouse",
            "communs-usage": "habitation", "communs-regime_juridique": "monopropriete", "communs-zone_tendue": "on",
            "bailleur-type": "physique", "bailleur-nom": "Garcia", "bailleur-prenom": "Ines",
            "lots-TOTAL_FORMS": "3", "lots-INITIAL_FORMS": "0", "lots-MIN_NUM_FORMS": "1", "lots-MAX_NUM_FORMS": "60",
            "lots-0-complement": "Apt 1", "lots-0-nb_pieces": "2", "lots-0-surface": "41.5",
            "lots-0-dernier_loyer": "620", "lots-0-charges": "40",
            "lots-1-numero_appartement": "2", "lots-1-etage": "1", "lots-1-meuble": "on",
            **autres,
        }

    def test_immeuble_et_ses_lots(self):
        """Un immeuble se crée avec ses lots : chacun devient une propriété."""
        self.assertContains(self.client.get(reverse("proprietes:propriete_creer")), reverse("proprietes:immeuble_creer"))
        self.assertContains(self.client.get(reverse("proprietes:immeuble_creer")), "+ Ajouter un lot")
        reponse = self.client.post(reverse("proprietes:immeuble_creer"), self.immeuble())
        immeuble = Immeuble.objects.get()
        self.assertRedirects(reponse, immeuble.get_absolute_url())
        premier, second = immeuble.lots.order_by("pk")
        self.assertEqual((premier.complement, premier.surface, premier.dernier_loyer), ("Apt 1", Decimal("41.5"), 620))
        self.assertEqual((second.numero_appartement, second.etage, second.meuble), ("2", "1", True))
        for lot in (premier, second):
            self.assertEqual((lot.adresse, lot.ville, lot.bailleur.nom), ("8 quai de Brienne", "Toulouse", "Garcia"))
            self.assertEqual((lot.type_habitat, lot.regime_juridique, lot.zone_tendue), ("collectif", "monopropriete", True))
        self.assertEqual(Bailleur.objects.count(), 1)
        page = self.client.get(immeuble.get_absolute_url())
        self.assertContains(page, "2 lots")
        self.assertContains(page, "Apt 1")
        # Chaque fiche dit ce que son bouton modifie (#160).
        self.assertContains(page, "Modifier l’immeuble")
        self.assertContains(self.client.get(reverse("proprietes:liste")), "Immeuble Résidence du Canal")
        self.assertContains(self.client.get(premier.get_absolute_url()), immeuble.get_absolute_url())
        self.assertContains(self.client.get(premier.get_absolute_url()), "Modifier le logement")
        # Le n° de lot désigne le lot dans la liste de l'immeuble.
        premier.numero_lot = "12"
        premier.save()
        self.assertContains(self.client.get(immeuble.get_absolute_url()), "Lot 12")
        # Un lot de plus reprend l'adresse et ce que les lots ont en commun.
        page = self.client.get(reverse("proprietes:propriete_creer"), {"immeuble": immeuble.pk})
        initial = page.context["form"].initial
        self.assertEqual((initial["adresse"], initial["bailleur"], initial["regime_juridique"]),
                         ("8 quai de Brienne", premier.bailleur_id, "monopropriete"))
        # Corriger l'adresse de l'immeuble corrige celle de ses lots.
        self.client.post(reverse("proprietes:immeuble_modifier", args=[immeuble.pk]), {
            "nom": "Résidence du Canal", "adresse": "8 bis quai de Brienne", "code_postal": "31000", "ville": "Toulouse",
        })
        self.assertEqual(set(immeuble.lots.values_list("adresse", flat=True)), {"8 bis quai de Brienne"})

    def test_mandat_d_abord(self):
        """« Nouveau mandat » est la porte d'entrée : la propriété se crée en
        chemin, puis son mandat s'ouvre prérempli."""
        page = self.client.get(reverse("mandats:creer"))
        self.assertContains(page, reverse("proprietes:propriete_creer") + "?suite=mandat")
        self.assertContains(page, reverse("proprietes:immeuble_creer") + "?suite=mandat")
        self.assertContains(self.client.get(reverse("proprietes:liste")), f'href="{reverse("mandats:creer")}">+ Nouvelle propriété')
        reponse = self.client.post(reverse("proprietes:propriete_creer") + "?suite=mandat", {
            "usage": "habitation", "actif": "on", "adresse": "3 place du Capitole", "code_postal": "31000",
            "ville": "Toulouse", "bailleur-type": "physique", "bailleur-nom": "Garcia", "bailleur-prenom": "Ines",
        })
        propriete = Propriete.objects.get()
        self.assertRedirects(reponse, f"{reverse('mandats:creer')}?propriete={propriete.pk}")
        # La liste dit d'un coup d'œil si la propriété a un mandat : aucun, à signer, ou son numéro.
        from mandats.models import Mandat

        self.assertContains(self.client.get(reverse("proprietes:liste")), "Sans mandat</span>")
        mandat = Mandat.depuis_propriete(propriete)
        mandat.loyer_cc, mandat.numero = 600, 412
        mandat.save()
        liste = self.client.get(reverse("proprietes:liste"))
        self.assertContains(liste, f'<a href="{mandat.get_absolute_url()}">n° 412</a>')
        self.assertContains(liste, "À signer</span>")
        self.assertNotContains(liste, "Sans mandat</span>")
        mandat.date_signature = dt.date(2026, 9, 1)
        mandat.save()
        self.assertNotContains(self.client.get(reverse("proprietes:liste")), "À signer</span>")
        mandat.delete()
        # Sans le parcours, la fiche reste enregistrable seule, et se retrouve par le filtre.
        self.assertContains(self.client.get(reverse("proprietes:liste"), {"mandat": "sans"}), "3 place du Capitole")
        # Un immeuble : un mandat par lot, depuis sa fiche.
        # La bailleresse, déjà enregistrée avec le logement, se choisit dans la liste.
        immeuble = self.immeuble(**{"communs-bailleur": propriete.bailleur_id})
        reponse = self.client.post(reverse("proprietes:immeuble_creer") + "?suite=mandat", immeuble, follow=True)
        self.assertContains(reponse, "Rédigez maintenant le mandat de chaque lot.")
        self.assertContains(reponse, "2 sans mandat")
        lot = Immeuble.objects.get().lots.first()
        self.assertContains(reponse, f"{reverse('mandats:creer')}?propriete={lot.pk}")

    def test_immeuble_incomplet(self):
        """Sans lot, ou avec un lot que rien ne distingue, rien n'est créé."""
        sans_lot = {cle: valeur for cle, valeur in self.immeuble().items() if not cle.startswith("lots-") or cle.endswith("FORMS")}
        self.assertEqual(self.client.post(reverse("proprietes:immeuble_creer"), sans_lot).status_code, 200)
        reponse = self.client.post(reverse("proprietes:immeuble_creer"), self.immeuble(**{"lots-1-numero_appartement": ""}))
        self.assertContains(reponse, "Indiquer une désignation ou un n° d&#x27;appartement.")
        self.assertFalse(Immeuble.objects.exists() or Propriete.objects.exists() or Bailleur.objects.exists())

    def test_documents_de_la_propriete(self):
        """Le DPE et les autres diagnostics se déposent sur la fiche."""
        propriete = Propriete.objects.create(
            bailleur=Bailleur.objects.create(nom="Durand"), adresse="1 rue", code_postal="31000", ville="Toulouse",
        )
        deposer = reverse("proprietes:piece_deposer", args=[propriete.pk])
        onglet = propriete.get_absolute_url() + "?onglet=documents"
        with tempfile.TemporaryDirectory() as dossier, self.settings(MEDIA_ROOT=dossier):
            self.assertContains(self.client.get(onglet), "Déposer un document")
            # Un fichier qui n'est ni un PDF ni une image est refusé, même renommé.
            faux = SimpleUploadedFile("dpe.pdf", b"<html>", content_type="application/pdf")
            reponse = self.client.post(deposer, {"type": "dpe", "fichier": faux}, follow=True)
            self.assertContains(reponse, "ni un PDF ni une image")
            self.assertFalse(Piece.objects.exists())
            dpe = SimpleUploadedFile("dpe.pdf", b"%PDF-1.7 essai", content_type="application/pdf")
            reponse = self.client.post(deposer, {"type": "dpe", "date_document": "2023-04-12", "fichier": dpe})
            self.assertRedirects(reponse, onglet)
            piece = Piece.objects.get()
            propriete.refresh_from_db()
            self.assertEqual(propriete.date_dpe, dt.date(2023, 4, 12))
            self.assertEqual((piece.expire_le, piece.perime, piece.depose_par.username), (dt.date(2033, 4, 12), False, "gestion"))
            self.assertContains(self.client.get(onglet), "jusqu'au 12/04/2033")
            self.assertContains(self.client.get(propriete.get_absolute_url() + "?onglet=informations"), "ouvrir le DPE")
            fichier = self.client.get(reverse("proprietes:piece", args=[piece.pk]))
            self.assertEqual((fichier["Content-Type"], fichier.getvalue()), ("application/pdf", b"%PDF-1.7 essai"))
            # Un document ne s'ouvre que connecté.
            self.client.logout()
            self.assertEqual(self.client.get(reverse("proprietes:piece", args=[piece.pk])).status_code, 302)
            self.client.force_login(User.objects.get(username="gestion"))
            self.client.post(reverse("proprietes:piece_supprimer", args=[piece.pk]))
            self.assertFalse(Piece.objects.exists())

    def test_validite_des_diagnostics(self):
        self.assertTrue(Piece(type="dpe", date_document=dt.date(2020, 5, 1)).perime)
        self.assertTrue(Piece(type="erp", date_document=dt.date(2020, 5, 1)).perime)
        self.assertIsNone(Piece(type="reglement", date_document=dt.date(2000, 1, 1)).expire_le)
        self.assertIsNone(Piece(type="dpe").expire_le)

    def test_societe_et_son_representant(self):
        """Pour une SCI, la raison sociale et la personne qui la représente
        sont deux contacts distincts ; la désignation des actes les réunit."""
        from contacts.models import Contact

        donnees = {
            "type": "sci_familiale", "nom": "SCI ESGT", "civilite": "Mme", "prenom": "Ignoré",
            "representant_civilite": "Mme", "representant_nom": "Garcia", "representant_prenom": "Elodie",
            "representant_qualite": "gérante", "siren": "732 829 320", "telephone": "0612345678",
            "associes-TOTAL_FORMS": "2", "associes-INITIAL_FORMS": "0",
            "associes-0-civilite": "Mme", "associes-0-nom": "Garcia", "associes-0-prenom": "Elodie", "associes-0-parts": "60",
            "associes-1-civilite": "M.", "associes-1-nom": "Garcia", "associes-1-prenom": "Tom", "associes-1-parts": "40",
        }
        self.assertContains(self.client.get(reverse("proprietes:bailleur_creer")), "js/bailleur.js")
        self.client.post(reverse("proprietes:bailleur_creer"), donnees)
        bailleur = Bailleur.objects.get()
        self.assertEqual((bailleur.prenom, bailleur.civilite, bailleur.telephone), ("", "", "06 12 34 56 78"))
        self.assertEqual(bailleur.designation, "SCI ESGT, représentée par Madame GARCIA Elodie, gérante")
        self.assertEqual(str(bailleur), "SCI ESGT")
        # Le SIREN mène à la fiche de la société, où son extrait RNE se télécharge.
        self.assertEqual(bailleur.lien_annuaire, "https://annuaire-entreprises.data.gouv.fr/entreprise/732829320")
        self.assertContains(self.client.get(bailleur.get_absolute_url()), bailleur.lien_annuaire)
        societe, personne = bailleur.contact, bailleur.representant
        self.assertEqual((societe.nature, societe.nom, societe.prenom), ("entreprise", "SCI ESGT", ""))
        self.assertEqual((personne.nature, personne.nom, personne.prenom, personne.civilite), ("personne", "Garcia", "Elodie", "Mme"))
        # Les associés ont leur fiche ; la gérante, aussi associée, n'en a qu'une.
        self.assertEqual(Contact.objects.count(), 3)
        self.assertEqual(dict(personne.etiquettes), {"representant": "Représentant", "associe": "Associé"})
        tom = Contact.objects.get(prenom="Tom")
        self.assertEqual(dict(tom.etiquettes), {"associe": "Associé"})
        self.assertIn("Associé de SCI ESGT", [lien.libelle for lien in tom.relations])
        # Chaque contact renvoie à l'autre : la gérante vers la société, la société vers sa gérante et son associé.
        self.assertEqual([(r.libelle, r.url) for r in personne.relations if r.type == "represente"],
                         [("Représente SCI ESGT", societe.get_absolute_url())])
        self.assertEqual([r.libelle for r in societe.relations], ["Représentée par GARCIA Elodie", "Associé : GARCIA Tom"])
        liste = self.client.get(reverse("contacts:liste"))
        self.assertContains(liste, f'<a href="{personne.get_absolute_url()}">Représentée par GARCIA Elodie</a>')
        self.assertContains(liste, f'<a href="{societe.get_absolute_url()}">Représente SCI ESGT</a>')
        # Le bloc dépliable des liens garde son clic : la ligne ne l'intercepte pas.
        self.assertContains(liste, 'closest("a, button, form, label, summary, details, input, select")')
        # La recherche par le nom de la société trouve aussi sa gérante.
        self.assertContains(self.client.get(reverse("contacts:liste"), {"q": "esgt"}), "GARCIA Elodie")
        self.assertContains(self.client.get(bailleur.get_absolute_url()), "GARCIA Tom</a> (40 %)")
        self.assertContains(self.client.get(societe.get_absolute_url()), "GARCIA Tom")
        self.assertContains(self.client.get(reverse("contacts:liste"), {"categorie": "associe"}), "GARCIA Tom")
        # Les parts ne dépassent pas 100 %.
        reponse = self.client.post(reverse("proprietes:bailleur_modifier", args=[bailleur.pk]), {
            **donnees, "associes-INITIAL_FORMS": "0", "associes-0-parts": "70",
        })
        self.assertContains(reponse, "elles ne peuvent dépasser 100 %")
        # La fiche contact d'une société ne porte ni civilité ni prénom, même si on en envoie.
        page = self.client.get(reverse("contacts:modifier", args=[societe.pk]))
        self.assertContains(page, "sa fiche de bailleur")
        self.client.post(reverse("contacts:modifier", args=[societe.pk]), {
            "nature": "entreprise", "nom": "SCI ESGT", "civilite": "Mme", "prenom": "Elodie", "date_naissance": "1980-01-01",
        })
        societe.refresh_from_db()
        self.assertEqual((societe.civilite, societe.prenom, societe.date_naissance), ("", "", None))
        self.assertContains(self.client.get(personne.get_absolute_url()), "SCI ESGT")
        self.assertContains(self.client.get(societe.get_absolute_url()), "GARCIA Elodie")
        self.assertContains(self.client.get(reverse("contacts:liste"), {"categorie": "representant"}), "GARCIA Elodie")
        # Changer de représentant relie la société à une autre personne.
        donnees = {cle: valeur for cle, valeur in donnees.items() if not cle.startswith("associes-")}
        donnees.update({"associes-TOTAL_FORMS": "0", "associes-INITIAL_FORMS": "0"})
        self.client.post(reverse("proprietes:bailleur_modifier", args=[bailleur.pk]), {**donnees, "representant_nom": "Faure", "representant_prenom": "Paul"})
        bailleur.refresh_from_db()
        self.assertEqual(bailleur.representant.nom, "Faure")
        # Une personne physique n'a pas de représentant.
        self.client.post(reverse("proprietes:bailleur_modifier", args=[bailleur.pk]), {**donnees, "type": "physique", "nom": "Garcia", "prenom": "Elodie"})
        bailleur.refresh_from_db()
        self.assertEqual((bailleur.representant, bailleur.representant_nom, bailleur.siren, bailleur.prenom), (None, "", "", "Elodie"))
        self.assertFalse(bailleur.associes.exists())

    def test_controles_de_saisie(self):
        """Téléphone, e-mail, code postal et SIREN sont contrôlés à la saisie."""
        from core.champs import normaliser_telephone

        for saisi, attendu in [("0612345678", "06 12 34 56 78"), ("06.12.34.56.78", "06 12 34 56 78"),
                               ("+33 6 12 34 56 78", "06 12 34 56 78"), ("0032 470 12 34 56", "+32470123456"), ("", "")]:
            self.assertEqual(normaliser_telephone(saisi), attendu)
        reponse = self.client.post(reverse("proprietes:bailleur_creer"), {
            "type": "morale", "nom": "Bernard SA", "siren": "123 456 789", "telephone": "06 12 34", "email": "contact@",
            "associes-TOTAL_FORMS": "0", "associes-INITIAL_FORMS": "0",
        })
        erreurs = reponse.context["form"].errors
        self.assertEqual(set(erreurs), {"siren", "telephone", "email"})
        self.assertFalse(Bailleur.objects.exists())
        bailleur = Bailleur.objects.create(nom="Durand")
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {
            "bailleur": bailleur.pk, "usage": "habitation", "adresse": "1 rue", "code_postal": "3100", "ville": "Toulouse",
        })
        self.assertContains(reponse, "Le code postal compte cinq chiffres.")

    def test_adresse_deja_enregistree(self):
        """Saisir une adresse déjà présente montre ce qui s'y trouve et
        demande confirmation ; un autre lot du même immeuble passe sans question."""
        bailleur = Bailleur.objects.create(nom="Durand")
        existante = Propriete.objects.create(bailleur=bailleur, adresse="302 avenue de Muret", code_postal="31300", ville="Toulouse")
        donnees = {"bailleur": bailleur.pk, "usage": "habitation", "actif": "on", "adresse": "302, Av. de MURET",
                   "code_postal": "31300", "ville": "Toulouse"}
        reponse = self.client.post(reverse("proprietes:propriete_creer"), donnees)
        self.assertContains(reponse, "Cette propriété semble déjà enregistrée.")
        self.assertContains(reponse, existante.get_absolute_url())
        self.assertEqual(Propriete.objects.count(), 1)
        # Un autre lot à la même adresse : signalé, puis accepté une fois confirmé.
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {**donnees, "complement": "Apt 2"})
        self.assertContains(reponse, "Cette adresse est déjà enregistrée.")
        self.client.post(reverse("proprietes:propriete_creer"), {**donnees, "complement": "Apt 2", "confirmer_adresse": "on"})
        self.assertEqual(Propriete.objects.count(), 2)
        # La même rue dans une autre commune n'est pas un doublon.
        self.client.post(reverse("proprietes:propriete_creer"), {**donnees, "code_postal": "31600", "ville": "Muret"})
        self.assertEqual(Propriete.objects.count(), 3)
        # Modifier une fiche sans toucher à son adresse ne pose pas la question.
        reponse = self.client.post(reverse("proprietes:propriete_modifier", args=[existante.pk]), {**donnees, "adresse": "302 avenue de Muret", "surface": "40"})
        self.assertRedirects(reponse, existante.get_absolute_url())
        # Un immeuble à une adresse déjà connue demande aussi confirmation ; ses lots ne se répètent pas.
        immeuble = self.immeuble(**{"immeuble-adresse": "302 avenue de Muret", "immeuble-code_postal": "31300"})
        reponse = self.client.post(reverse("proprietes:immeuble_creer"), immeuble)
        self.assertContains(reponse, "Créer quand même cet immeuble et ses lots")
        self.assertFalse(Immeuble.objects.exists())
        reponse = self.client.post(reverse("proprietes:immeuble_creer"), {
            **self.immeuble(), "lots-1-numero_appartement": "", "lots-1-complement": "APT 1",
        })
        self.assertContains(reponse, "Ce lot figure déjà dans une autre ligne.")
        self.client.post(reverse("proprietes:immeuble_creer"), self.immeuble())
        lot = Immeuble.objects.get().lots.first()
        # Ajouter un lot distinct à cet immeuble passe sans confirmation.
        self.client.post(reverse("proprietes:propriete_creer") + f"?immeuble={lot.immeuble_id}", {
            "bailleur": lot.bailleur_id, "immeuble": lot.immeuble_id, "usage": "habitation", "actif": "on",
            "adresse": lot.adresse, "code_postal": lot.code_postal, "ville": lot.ville, "complement": "Apt 9",
        })
        self.assertEqual(Immeuble.objects.get().lots.count(), 3)

    def test_modifier_et_supprimer_depuis_la_fiche(self):
        from mandats.models import Mandat

        bailleur = Bailleur.objects.create(nom="Durand")
        erreur = Propriete.objects.create(bailleur=bailleur, adresse="1 rue Erreur", code_postal="31000", ville="Toulouse")
        geree = Propriete.objects.create(
            bailleur=bailleur, adresse="2 rue Gérée", code_postal="31000", ville="Toulouse", dernier_loyer=600,
        )
        Mandat.depuis_propriete(geree).save()
        # La liste ouvre la fiche ; « Modifier » et « Supprimer » sont en tête de la fiche.
        liste = self.client.get(reverse("proprietes:liste"))
        self.assertNotContains(liste, reverse("proprietes:propriete_supprimer", args=[erreur.pk]))
        page = self.client.get(erreur.get_absolute_url())
        self.assertContains(page, '<nav class="fil"')
        self.assertContains(page, reverse("proprietes:propriete_modifier", args=[erreur.pk]))
        self.assertContains(page, reverse("proprietes:propriete_supprimer", args=[erreur.pk]))
        # Une propriété sous mandat ne se supprime pas.
        reponse = self.client.post(reverse("proprietes:propriete_supprimer", args=[geree.pk]), follow=True)
        self.assertContains(reponse, "Suppression refusée : cette propriété a 1 mandat (le mandat n° 1 est en cours)")
        self.assertTrue(Propriete.objects.filter(pk=geree.pk).exists())
        reponse = self.client.post(reverse("proprietes:propriete_supprimer", args=[erreur.pk]), follow=True)
        self.assertContains(reponse, "Propriété supprimée : 1 rue Erreur")
        self.assertEqual(list(Propriete.objects.all()), [geree])
        self.assertEqual(self.client.get(reverse("proprietes:propriete_supprimer", args=[geree.pk])).status_code, 405)

    def test_supprimer_bailleur_immeuble_et_contact(self):
        from contacts.models import Contact

        # Un immeuble se supprime sans emporter ses lots.
        self.client.post(reverse("proprietes:immeuble_creer"), self.immeuble())
        immeuble = Immeuble.objects.get()
        self.assertContains(self.client.get(immeuble.get_absolute_url()), reverse("proprietes:immeuble_supprimer", args=[immeuble.pk]))
        reponse = self.client.post(reverse("proprietes:immeuble_supprimer", args=[immeuble.pk]), follow=True)
        self.assertContains(reponse, "Ses 2 lots sont conservés")
        self.assertEqual((Immeuble.objects.count(), Propriete.objects.filter(immeuble__isnull=True).count()), (0, 2))
        # Un bailleur qui a des propriétés reste ; sans propriété, il se supprime.
        bailleur = Bailleur.objects.get()
        reponse = self.client.post(reverse("proprietes:bailleur_supprimer", args=[bailleur.pk]), follow=True)
        self.assertContains(reponse, "Suppression refusée : ce bailleur a 2 propriétés.")
        contact = bailleur.contact
        reponse = self.client.post(reverse("contacts:supprimer", args=[contact.pk]), follow=True)
        self.assertContains(reponse, "Suppression refusée : ce contact est relié")
        Propriete.objects.all().delete()
        self.client.post(reverse("proprietes:bailleur_supprimer", args=[bailleur.pk]))
        self.assertFalse(Bailleur.objects.exists())
        # Le contact, que plus rien ne relie, se supprime à son tour.
        self.assertContains(self.client.get(contact.get_absolute_url()), reverse("contacts:supprimer", args=[contact.pk]))
        self.client.post(reverse("contacts:supprimer", args=[contact.pk]))
        self.assertFalse(Contact.objects.filter(pk=contact.pk).exists())
        self.assertEqual(self.client.get(reverse("contacts:supprimer", args=[contact.pk])).status_code, 405)

    def test_import_hors_session_et_export_sans_formule(self):
        """Le fichier vérifié attend sa confirmation sur le disque, pas dans la
        session ; un texte qui ressemble à une formule est neutralisé à l'export."""
        with tempfile.TemporaryDirectory() as dossier, self.settings(MEDIA_ROOT=dossier):
            fichier = SimpleUploadedFile("proprietes.xlsx", xlsx(ligne("L001", **{"Complément": "@SUM(1+1)"})))
            self.client.post(reverse("proprietes:importer"), {"fichier": fichier})
            depot = self.client.session["import_proprietes"]
            self.assertNotIn("contenu", depot)
            self.assertRegex(depot["depot"], r"^[0-9a-f]{32}$")
            self.client.post(reverse("proprietes:importer_confirmer"))
            self.assertEqual(Propriete.objects.count(), 1)
            # Le fichier temporaire est effacé, et un nom de dépôt forgé est ignoré.
            import pathlib
            self.assertEqual(list((pathlib.Path(dossier) / "imports").iterdir()), [])
            session = self.client.session
            session["import_proprietes"] = {"nom": "x.xlsx", "depot": "../../etc/passwd"}
            session.save()
            reponse = self.client.post(reverse("proprietes:importer_confirmer"), follow=True)
            self.assertContains(reponse, "Aucun fichier en attente")
        classeur = load_workbook(io.BytesIO(self.client.get(reverse("proprietes:exporter")).content))
        valeurs = [cellule.value for cellule in classeur.active[2]]
        self.assertIn("'@SUM(1+1)", valeurs)
        self.assertEqual(tableur.sans_formule("+33 6 12 34 56 78"), "'+33 6 12 34 56 78")
        self.assertEqual(tableur.sans_formule("12 rue des Lois"), "12 rue des Lois")
        # Un fichier démesuré est refusé avant d'être traité.
        with self.assertRaisesMessage(ValueError, "Fichier trop long"):
            tableur.lire_lignes("gros.csv", ("a;b\n" * (tableur.LIGNES_MAX + 2)).encode())

    def test_un_immeuble_une_seule_fois(self):
        """Le champ « Immeuble » d'une propriété propose chaque immeuble une
        fois, avec son adresse ; le même immeuble ne s'enregistre pas deux fois."""
        from django.db import IntegrityError, transaction

        bailleur = Bailleur.objects.create(nom="SCI ESGT", type="sci_familiale")
        mermoz = Immeuble.objects.create(nom="MERMOZ", adresse="2 Rue Jean Mermoz", code_postal="31120", ville="Portet-sur-Garonne")
        canal = Immeuble.objects.create(nom="Résidence du Canal", adresse="8 quai de Brienne", code_postal="31000", ville="Toulouse")
        sans_nom = Immeuble.objects.create(adresse="5 rue du Taur", code_postal="31000", ville="Toulouse")
        lots = [
            Propriete.objects.create(bailleur=bailleur, immeuble=mermoz, adresse=mermoz.adresse, code_postal="31120",
                                     ville=mermoz.ville, complement=f"Apt{numero}")
            for numero in (1, 2, 3)
        ]
        page = self.client.get(reverse("proprietes:propriete_modifier", args=[lots[0].pk]))
        libelles = [str(libelle) for valeur, libelle in page.context["form"].fields["immeuble"].choices if valeur]
        # Trois lots, un seul choix « MERMOZ » ; les autres immeubles sont là, chacun une fois.
        self.assertEqual(libelles, [
            "5 rue du Taur, 31000 Toulouse",
            "MERMOZ — 2 Rue Jean Mermoz, 31120 Portet-sur-Garonne",
            "Résidence du Canal — 8 quai de Brienne, 31000 Toulouse",
        ])
        self.assertContains(page, f'<option value="{mermoz.pk}" selected>MERMOZ — 2 Rue Jean Mermoz, 31120 Portet-sur-Garonne</option>', html=True)
        # Enregistrer la fiche garde le rattachement et ne crée ni immeuble ni propriété.
        donnees = {"bailleur": bailleur.pk, "immeuble": mermoz.pk, "usage": "habitation", "actif": "on",
                   "adresse": mermoz.adresse, "code_postal": "31120", "ville": mermoz.ville, "complement": "Apt1", "surface": "42"}
        reponse = self.client.post(reverse("proprietes:propriete_modifier", args=[lots[0].pk]), donnees)
        self.assertRedirects(reponse, lots[0].get_absolute_url())
        lots[0].refresh_from_db()
        self.assertEqual((lots[0].immeuble, lots[0].surface), (mermoz, Decimal("42")))
        self.assertEqual((Immeuble.objects.count(), Propriete.objects.count(), mermoz.lots.count()), (3, 3, 3))
        # Le même immeuble, écrit autrement, est refusé à la création…
        memes = self.immeuble(**{"immeuble-nom": "Mermoz", "immeuble-adresse": "2, rue Jean MERMOZ",
                                 "immeuble-code_postal": "31120", "immeuble-ville": "Portet sur Garonne", "immeuble-confirmer_adresse": "on"})
        reponse = self.client.post(reverse("proprietes:immeuble_creer"), memes)
        self.assertContains(reponse, "Cet immeuble est déjà enregistré")
        self.assertContains(reponse, mermoz.get_absolute_url())
        self.assertEqual((Immeuble.objects.count(), Propriete.objects.count()), (3, 3))
        # … en renommant un autre immeuble…
        reponse = self.client.post(reverse("proprietes:immeuble_modifier", args=[canal.pk]), {
            "nom": "MERMOZ", "adresse": "2 rue Jean Mermoz", "code_postal": "31120", "ville": "Portet-sur-Garonne",
        })
        self.assertContains(reponse, "Cet immeuble est déjà enregistré")
        # … et par la base elle-même.
        with self.assertRaises(IntegrityError), transaction.atomic():
            Immeuble.objects.create(nom="mermoz", adresse="2 RUE JEAN MERMOZ", code_postal="31120", ville="Portet-sur-Garonne")
        # Un autre bâtiment à la même adresse reste possible sous un autre nom, après confirmation.
        autre = {**memes, "immeuble-nom": "MERMOZ bâtiment B"}
        self.client.post(reverse("proprietes:immeuble_creer"), autre)
        self.assertEqual(Immeuble.objects.count(), 4)
        # Modifier un immeuble sans changer son identité ne le prend pas pour son propre doublon.
        reponse = self.client.post(reverse("proprietes:immeuble_modifier", args=[sans_nom.pk]), {
            "nom": "", "adresse": "5 rue du Taur", "code_postal": "31000", "ville": "Toulouse",
        })
        self.assertRedirects(reponse, sans_nom.get_absolute_url())

    def test_un_bailleur_une_seule_fiche(self):
        """Saisir à nouveau un bailleur déjà connu est refusé : on reprend sa fiche."""
        vide = {"associes-TOTAL_FORMS": "0", "associes-INITIAL_FORMS": "0"}
        sci = {"type": "sci_familiale", "nom": "SCI ESGT", "representant_nom": "Couet", "representant_prenom": "Elodie", **vide}
        self.client.post(reverse("proprietes:bailleur_creer"), sci)
        reponse = self.client.post(reverse("proprietes:bailleur_creer"), {**sci, "nom": "sci  esgt", "representant_qualite": "gérante"})
        self.assertContains(reponse, "Ce bailleur est déjà enregistré")
        self.assertEqual(Bailleur.objects.count(), 1)
        # Depuis la création d'une propriété aussi : le bailleur se choisit dans la liste.
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {
            "usage": "habitation", "actif": "on", "adresse": "2 rue Jean Mermoz", "code_postal": "31120",
            "ville": "Portet-sur-Garonne", "bailleur-type": "sci_familiale", "bailleur-nom": "SCI ESGT",
        })
        self.assertContains(reponse, "dans la liste « Bailleur » ci-dessus")
        self.assertEqual((Bailleur.objects.count(), Propriete.objects.count()), (1, 0))
        # Deux sociétés de même nom mais de SIREN différents sont bien deux bailleurs ; deux homonymes
        # se distinguent par leur code postal.
        existant = Bailleur.objects.get()
        existant.siren = "732 829 320"
        existant.save()
        self.client.post(reverse("proprietes:bailleur_creer"), {**sci, "siren": "552 100 554"})
        self.assertEqual(Bailleur.objects.count(), 2)
        personne = {"type": "physique", "nom": "Durand", "prenom": "Paul", "code_postal": "31000", **vide}
        self.client.post(reverse("proprietes:bailleur_creer"), personne)
        self.assertContains(self.client.post(reverse("proprietes:bailleur_creer"), {**personne, "nom": "DURAND"}), "déjà enregistré")
        self.client.post(reverse("proprietes:bailleur_creer"), {**personne, "code_postal": "75011"})
        self.assertEqual(Bailleur.objects.filter(nom="Durand").count(), 2)
        # Modifier une fiche ne la prend pas pour son propre doublon.
        reponse = self.client.post(reverse("proprietes:bailleur_modifier", args=[existant.pk]), {**sci, "siren": "732 829 320", "email": "a@b.fr"})
        self.assertRedirects(reponse, existant.get_absolute_url())

    def test_depenses_energie_incoherentes(self):
        bailleur = Bailleur.objects.create(nom="Durand")
        reponse = self.client.post(reverse("proprietes:propriete_creer"), {
            "bailleur": bailleur.pk, "usage": "habitation", "adresse": "1 rue",
            "code_postal": "31000", "ville": "Toulouse",
            "depenses_energie_min": 900, "depenses_energie_max": 500,
        })
        self.assertContains(reponse, "Le maximum doit être supérieur au minimum")

    def test_toutes_les_pages_s_affichent(self):
        tableur.importer("proprietes.xlsx", xlsx(ligne("L001")))
        propriete, bailleur = Propriete.objects.get(), Bailleur.objects.get()
        for url in [
            reverse("proprietes:liste"), reverse("proprietes:propriete", args=[propriete.pk]),
            reverse("proprietes:propriete_creer"), reverse("proprietes:propriete_modifier", args=[propriete.pk]),
            reverse("contacts:liste"), reverse("proprietes:bailleur", args=[bailleur.pk]),
            reverse("proprietes:bailleur_creer"), reverse("proprietes:bailleur_modifier", args=[bailleur.pk]),
            reverse("proprietes:importer"), reverse("proprietes:modele"), reverse("accueil"),
        ]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
